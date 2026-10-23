# routers/insurance.py
from fastapi import APIRouter, Depends, Request, Form
from fastapi.responses import HTMLResponse, RedirectResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy.orm import Session
from database import get_db
from models.insurance import InsurancePolicy, PremiumPayment, InsuranceNominee, InsuranceValuation
from models.persons import Person
from models.enums import InsuranceType, InsuranceStatus, PaymentMethod, PaymentStatus
from decimal import Decimal
from datetime import date, timedelta
from dateutil.relativedelta import relativedelta
from sqlalchemy import func

router = APIRouter(prefix="/insurance", tags=["insurance"])
templates = Jinja2Templates(directory="templates")

# Register custom format_inr filter for Indian numbering system
from routers.mutual_funds import format_inr
templates.env.filters["format_inr"] = format_inr

def format_lakhs(value) -> str:
    if value is None:
        return "—"
    try:
        val_float = float(value)
        if val_float == 0:
            return "₹0L"
        lakhs = val_float / 100000.0
        if lakhs.is_integer():
            return f"₹{int(lakhs)}L"
        else:
            formatted = f"{lakhs:.2f}"
            if formatted.endswith(".00"):
                formatted = formatted[:-3]
            elif formatted.endswith("0") and "." in formatted:
                formatted = formatted[:-1]
            return f"₹{formatted}L"
    except (ValueError, TypeError):
        return "—"

templates.env.filters["format_lakhs"] = format_lakhs


def clean_insurer_name(name: str) -> str:
    if not name:
        return "—"
    name_lower = name.lower()
    if "bharti" in name_lower and "axa" in name_lower:
        return "Bharti Axa"
    if "axis" in name_lower or "max" in name_lower:
        return "Axis Max"
    if "lic" in name_lower:
        return "LIC"
    if "kotak" in name_lower:
        return "Kotak"
    if "hdfc" in name_lower:
        return "HDFC"
    if "bajaj" in name_lower:
        return "Bajaj Allianz"
    if "tata" in name_lower:
        return "Tata AIA"
    if "canara" in name_lower:
        return "Canara HSBC"
    if "bandhan" in name_lower:
        return "Bandhan"
    
    for suffix in ["life insurance", "insurance", "corporation"]:
        if suffix in name_lower:
            idx = name_lower.find(suffix)
            return name[:idx].strip()
    return name.strip()


# ---------- POLICY LIST ----------
@router.get("/", response_class=HTMLResponse)
def list_policies(request: Request, db: Session = Depends(get_db)):
    policies = db.query(InsurancePolicy).join(
        Person, InsurancePolicy.insured_person_id == Person.id
    ).add_columns(Person.short_name).all()
    
    from services.performance_engine import calculate_xirr
    
    enriched = []
    for policy, short_name in policies:
        # Latest valuation fund value
        latest_val = db.query(InsuranceValuation).filter(
            InsuranceValuation.policy_id == policy.id
        ).order_by(InsuranceValuation.valuation_date.desc()).first()
        
        fund_value = None
        if latest_val:
            fund_value = latest_val.fund_value or latest_val.total_value or latest_val.surrender_value
        
        # XIRR calculation
        policy_xirr = None
        paid_payments = db.query(PremiumPayment).filter(
            PremiumPayment.policy_id == policy.id,
            PremiumPayment.status == PaymentStatus.PAID
        ).order_by(PremiumPayment.payment_date).all()
        
        if paid_payments:
            dates = []
            flows = []
            has_terminal_payout = False
            for p in paid_payments:
                txn_type = getattr(p, 'transaction_type', None) or 'PREMIUM'
                if txn_type in {'PARTIAL_WITHDRAWAL', 'SURRENDER', 'MATURITY_PAYOUT'}:
                    flows.append(float(p.amount))
                    if txn_type in {'SURRENDER', 'MATURITY_PAYOUT'}:
                        has_terminal_payout = True
                else:
                    flows.append(-float(p.amount))
                dates.append(p.payment_date)
            
            if not has_terminal_payout:
                inflow = None
                inflow_date = None
                if policy.status.value == "SURRENDERED" and policy.surrendered_amount is not None:
                    inflow = float(policy.surrendered_amount)
                    inflow_date = latest_val.valuation_date if latest_val else date.today()
                elif policy.status.value == "MATURED" and policy.maturity_amount is not None:
                    inflow = float(policy.maturity_amount)
                    inflow_date = policy.policy_end_date or date.today()
                elif fund_value is not None:
                    inflow = float(fund_value)
                    inflow_date = latest_val.valuation_date if latest_val else date.today()
                if inflow is not None and inflow > 0:
                    dates.append(inflow_date)
                    flows.append(inflow)
            
            try:
                xirr_val = calculate_xirr(dates, flows)
                if xirr_val is not None:
                    policy_xirr = xirr_val * 100
            except Exception:
                pass
        
        # Policy Age
        policy_age = "—"
        if policy.policy_start_date:
            today = date.today()
            diff = relativedelta(today, policy.policy_start_date)
            years = diff.years + diff.months / 12.0
            if years >= 1:
                policy_age = f"{years:.1f} Yr"
            else:
                months = diff.months + diff.days / 30.0
                policy_age = f"{months:.1f} Mo"

        # Premium Paid
        total_premium_to_be_paid = (policy.premium_amount or 0) * (policy.premium_term or 0)
        premium_paid_till_date = db.query(func.sum(PremiumPayment.amount)).filter(
            PremiumPayment.policy_id == policy.id,
            PremiumPayment.status == PaymentStatus.PAID,
            (PremiumPayment.transaction_type == "PREMIUM") | (PremiumPayment.transaction_type == None)
        ).scalar() or 0
        
        paid_premium_count = db.query(func.count(PremiumPayment.id)).filter(
            PremiumPayment.policy_id == policy.id,
            PremiumPayment.status == PaymentStatus.PAID,
            (PremiumPayment.transaction_type == "PREMIUM") | (PremiumPayment.transaction_type == None)
        ).scalar() or 0
        
        expected_payments = policy.premium_term or 0
        if policy.premium_frequency == "MONTHLY":
            expected_payments = (policy.premium_term or 0) * 12
        elif policy.premium_frequency == "QUARTERLY":
            expected_payments = (policy.premium_term or 0) * 4
        elif policy.premium_frequency == "HALF_YEARLY":
            expected_payments = (policy.premium_term or 0) * 2
        elif policy.premium_frequency == "SINGLE":
            expected_payments = 1
            
        if expected_payments > 0 and paid_premium_count >= expected_payments:
            premium_paid_till_date = total_premium_to_be_paid

        # Lock Status
        lock_status = "—"
        if policy.policy_start_date and policy.premium_term:
            lock_release_date = policy.policy_start_date + relativedelta(years=policy.premium_term + 1)
            if date.today() >= lock_release_date:
                lock_status = "UNLOCKED"
            else:
                lock_status = "LOCKED"

        # Available Withdrawal (Fund Value - 1.5 * Premium Amount)
        # For policies which are locked, keep withdrawable as blank (None)
        withdrawable = None
        if lock_status != "LOCKED":
            if fund_value is not None and policy.premium_amount is not None:
                calc_val = float(fund_value) - 1.5 * float(policy.premium_amount)
                withdrawable = max(0.0, calc_val)

        enriched.append({
            "policy": policy,
            "short_name": short_name,
            "fund_value": fund_value,
            "xirr": policy_xirr,
            "policy_age": policy_age,
            "premium_paid": premium_paid_till_date,
            "lock_status": lock_status,
            "withdrawable": withdrawable,
            "clean_insurer": clean_insurer_name(policy.insurer_name),
        })
    
    # Sort stable-alphabetically by insurer name to match Excel grouping
    enriched.sort(key=lambda x: x["clean_insurer"])
    
    return templates.TemplateResponse(
        request=request,
        name="insurance/list.html",
        context={"policies": enriched}
    )


# ---------- EXPORT TO EXCEL ----------
@router.get("/export")
def export_policies_excel(db: Session = Depends(get_db)):
    import io
    from fastapi.responses import StreamingResponse
    import pandas as pd
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    policies = db.query(InsurancePolicy).join(
        Person, InsurancePolicy.insured_person_id == Person.id
    ).add_columns(Person.short_name).all()
    
    from services.performance_engine import calculate_xirr
    
    rows = []
    for policy, short_name in policies:
        latest_val = db.query(InsuranceValuation).filter(
            InsuranceValuation.policy_id == policy.id
        ).order_by(InsuranceValuation.valuation_date.desc()).first()
        
        fund_value = None
        if latest_val:
            fund_value = latest_val.fund_value or latest_val.total_value or latest_val.surrender_value
            
        policy_xirr = None
        paid_payments = db.query(PremiumPayment).filter(
            PremiumPayment.policy_id == policy.id,
            PremiumPayment.status == PaymentStatus.PAID
        ).order_by(PremiumPayment.payment_date).all()
        
        if paid_payments:
            dates = []
            flows = []
            has_terminal_payout = False
            for p in paid_payments:
                txn_type = getattr(p, 'transaction_type', None) or 'PREMIUM'
                if txn_type in {'PARTIAL_WITHDRAWAL', 'SURRENDER', 'MATURITY_PAYOUT'}:
                    flows.append(float(p.amount))
                    if txn_type in {'SURRENDER', 'MATURITY_PAYOUT'}:
                        has_terminal_payout = True
                else:
                    flows.append(-float(p.amount))
                dates.append(p.payment_date)
            
            if not has_terminal_payout:
                inflow = None
                inflow_date = None
                if policy.status.value == "SURRENDERED" and policy.surrendered_amount is not None:
                    inflow = float(policy.surrendered_amount)
                    inflow_date = latest_val.valuation_date if latest_val else date.today()
                elif policy.status.value == "MATURED" and policy.maturity_amount is not None:
                    inflow = float(policy.maturity_amount)
                    inflow_date = policy.policy_end_date or date.today()
                elif fund_value is not None:
                    inflow = float(fund_value)
                    inflow_date = latest_val.valuation_date if latest_val else date.today()
                if inflow is not None and inflow > 0:
                    dates.append(inflow_date)
                    flows.append(inflow)
            
            try:
                xirr_val = calculate_xirr(dates, flows)
                if xirr_val is not None:
                    policy_xirr = xirr_val * 100
            except Exception:
                pass
                
        # Policy Age
        policy_age = None
        if policy.policy_start_date:
            today = date.today()
            diff = relativedelta(today, policy.policy_start_date)
            policy_age = diff.years + diff.months / 12.0

        # Premium Paid
        total_premium_to_be_paid = (policy.premium_amount or 0) * (policy.premium_term or 0)
        premium_paid_till_date = db.query(func.sum(PremiumPayment.amount)).filter(
            PremiumPayment.policy_id == policy.id,
            PremiumPayment.status == PaymentStatus.PAID,
            (PremiumPayment.transaction_type == "PREMIUM") | (PremiumPayment.transaction_type == None)
        ).scalar() or 0
        
        paid_premium_count = db.query(func.count(PremiumPayment.id)).filter(
            PremiumPayment.policy_id == policy.id,
            PremiumPayment.status == PaymentStatus.PAID,
            (PremiumPayment.transaction_type == "PREMIUM") | (PremiumPayment.transaction_type == None)
        ).scalar() or 0
        
        expected_payments = policy.premium_term or 0
        if policy.premium_frequency == "MONTHLY":
            expected_payments = (policy.premium_term or 0) * 12
        elif policy.premium_frequency == "QUARTERLY":
            expected_payments = (policy.premium_term or 0) * 4
        elif policy.premium_frequency == "HALF_YEARLY":
            expected_payments = (policy.premium_term or 0) * 2
        elif policy.premium_frequency == "SINGLE":
            expected_payments = 1
            
        if expected_payments > 0 and paid_premium_count >= expected_payments:
            premium_paid_till_date = total_premium_to_be_paid

        # Lock Status
        lock_status = "—"
        if policy.policy_start_date and policy.premium_term:
            lock_release_date = policy.policy_start_date + relativedelta(years=policy.premium_term + 1)
            if date.today() >= lock_release_date:
                lock_status = "UNLOCKED"
            else:
                lock_status = "LOCKED"

        # Available Withdrawal
        # For policies which are locked, keep withdrawable as blank (None)
        withdrawable = None
        if lock_status != "LOCKED":
            if fund_value is not None and policy.premium_amount is not None:
                calc_val = float(fund_value) - 1.5 * float(policy.premium_amount)
                withdrawable = max(0.0, calc_val)

        # Holder Name
        holder = db.query(Person).filter(Person.id == policy.holder_person_id).first()
        holder_name = holder.full_name if holder else "—"

        rows.append({
            "Insured": getattr(policy, "insured_person_name", None) or short_name,
            "Insurer": clean_insurer_name(policy.insurer_name),
            "Type": policy.insurance_type.value if policy.insurance_type else "",
            "Sum Assured\n(Lakhs)": float(policy.sum_assured) / 100000.0 if policy.sum_assured else None,
            "Start Date": policy.policy_start_date.isoformat() if policy.policy_start_date else "",
            "End Date": policy.policy_end_date.isoformat() if policy.policy_end_date else "",
            "Age\n(Years)": round(policy_age, 1) if policy_age is not None else None,
            "PPT\n(Years)": policy.premium_term or 0,
            "Paid Premium\n(Lakhs)": float(premium_paid_till_date) / 100000.0 if premium_paid_till_date else 0.0,
            "Func Value\n(Lakhs)": float(fund_value) / 100000.0 if fund_value is not None else None,
            "XIRR (%)": float(policy_xirr) / 100.0 if policy_xirr is not None else None,
            "Status": policy.status.value if policy.status else "",
            "Withdrawable\n(Lakhs)": float(withdrawable) / 100000.0 if withdrawable is not None else None
        })
        
    # Sort stable-alphabetically by insurer name to group them nicely
    rows.sort(key=lambda r: r["Insurer"])
    df = pd.DataFrame(rows)
    
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Policies')
        workbook = writer.book
        worksheet = writer.sheets['Policies']
        
        # Row height for header
        worksheet.row_dimensions[1].height = 40
        
        header_font = Font(name='Segoe UI', size=11, bold=True, color='000000')
        header_fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid') # Solid Yellow
        header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        
        center_align = Alignment(horizontal='center', vertical='center')
        right_align = Alignment(horizontal='right', vertical='center')
        left_align = Alignment(horizontal='left', vertical='center')
        
        for col_idx in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )
        
        # Alternating row fill (light grey)
        zebra_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
        
        # Alignments mapping by column letter
        alignments = {
            'A': left_align,    # Insured
            'B': left_align,    # Insurer
            'C': center_align,  # Type
            'D': right_align,   # Sum Assured (Lakhs)
            'E': center_align,  # Start Date
            'F': center_align,  # End Date
            'G': center_align,  # Age (Years)
            'H': center_align,  # PPT (Years)
            'I': right_align,   # Paid Premium (Lakhs)
            'J': right_align,   # Func Value (Lakhs)
            'K': right_align,   # XIRR (%)
            'L': center_align,  # Status
            'M': right_align    # Withdrawable (Lakhs)
        }
        
        col_widths = {
            'A': 8,   # Insured
            'B': 10,  # Insurer
            'C': 7,   # Type
            'D': 8,   # Sum Assured (Lakhs)
            'E': 11,  # Start Date
            'F': 11,  # End Date
            'G': 6,   # Age (Years)
            'H': 6,   # PPT (Years)
            'I': 8,   # Paid Premium (Lakhs)
            'J': 8,   # Func Value (Lakhs)
            'K': 7,   # XIRR (%)
            'L': 8,   # Status
            'M': 9    # Withdrawable (Lakhs)
        }
        
        # Set print view properties
        worksheet.page_setup.orientation = worksheet.ORIENTATION_LANDSCAPE
        worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
        
        # Narrow margins (0.25 inches left/right, 0.5 inches top/bottom)
        worksheet.page_margins.left = 0.25
        worksheet.page_margins.right = 0.25
        worksheet.page_margins.top = 0.5
        worksheet.page_margins.bottom = 0.5
        worksheet.page_margins.header = 0.2
        worksheet.page_margins.footer = 0.2
        
        # Enable AutoFilter on the table headers (A1:M{num_rows})
        worksheet.auto_filter.ref = f"A1:M{len(df) + 1}"
        
        # Show gridlines when printing and viewing
        worksheet.print_options.gridLines = True
        if hasattr(worksheet, 'views') and worksheet.views.sheetView:
            worksheet.views.sheetView[0].showGridLines = True
            
        for col in worksheet.columns:
            col_letter = get_column_letter(col[0].column)
            
            # Format dimensions
            worksheet.column_dimensions[col_letter].width = col_widths.get(col_letter, 12)
            
            for cell in col:
                cell.border = thin_border
                cell.font = Font(name='Segoe UI', size=10)
                
                if cell.row > 1:
                    cell.alignment = alignments.get(col_letter, left_align)
                    if cell.row % 2 == 1:
                        cell.fill = zebra_fill
                    
                    # Spacing for row height
                    worksheet.row_dimensions[cell.row].height = 20
                    
                    if isinstance(cell.value, (int, float)):
                        if col_letter in ['D', 'I', 'J', 'M']: # Financial (Lakhs)
                            cell.number_format = '#,##0.00'
                        elif col_letter == 'G': # Age (Years)
                            cell.number_format = '0.0'
                        elif col_letter == 'K': # XIRR (%)
                            cell.number_format = '0.00%'
            
    output.seek(0)
    
    headers = {
        'Content-Disposition': 'attachment; filename="insurance_policies.xlsx"'
    }
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers
    )


# ---------- ADD POLICY ----------
@router.get("/add", response_class=HTMLResponse)
def add_policy_form(request: Request, db: Session = Depends(get_db)):
    persons = db.query(Person).all()
    insurers = [
        "Bharti Axa Life Insurance",
        "Axis Max Life Insurance",
        "LIC (Life Insurance Corporation)",
        "Kotak Life Insurance",
        "HDFC Life Insurance",
        "Bajaj Allianz Life Insurance",
        "Tata AIA Life Insurance",
    ]
    return templates.TemplateResponse(
        request=request,
        name="insurance/add_policy.html",
        context={
            "persons": persons,
            "insurance_types": [e.value for e in InsuranceType],
            "statuses": [e.value for e in InsuranceStatus],
            "payment_methods": [e.value for e in PaymentMethod],
            "insurers": insurers,
        }
    )

def generate_auto_payments(db: Session, policy: InsurancePolicy):
    # Automatic Generation of Payments till today (respecting term limit)
    if policy.premium_amount and policy.premium_frequency and policy.policy_start_date:
        today = date.today()
        current_date = policy.policy_start_date
        
        # Calculate amount per payment based on "premium_amount" being "per year"
        p_amount = policy.premium_amount
        if policy.premium_frequency == "MONTHLY":
            p_amount = p_amount / Decimal(12)
        elif policy.premium_frequency == "QUARTERLY":
            p_amount = p_amount / Decimal(4)
        elif policy.premium_frequency == "HALF_YEARLY":
            p_amount = p_amount / Decimal(2)

        # We stop either at today OR at the end of the premium term
        stop_date = today
        if policy.premium_end_date:
            stop_date = min(today, policy.premium_end_date)

        while current_date < stop_date:
            payment = PremiumPayment(
                policy_id=policy.id,
                payment_date=current_date,
                amount=p_amount,
                payment_mode=policy.payment_method.value if policy.payment_method else "AUTO_GENERATED",
                status=PaymentStatus.PAID,
                notes="Automatically generated on policy creation"
            )
            db.add(payment)
            
            # Advance current_date
            if policy.premium_frequency == "MONTHLY":
                current_date += relativedelta(months=1)
            elif policy.premium_frequency == "QUARTERLY":
                current_date += relativedelta(months=3)
            elif policy.premium_frequency == "HALF_YEARLY":
                current_date += relativedelta(months=6)
            elif policy.premium_frequency == "YEARLY" or policy.premium_frequency == "ANNUAL":
                current_date += relativedelta(years=1)
            else:
                break

@router.post("/add")
def add_policy(
    policy_number: str = Form(...),
    insurance_type: str = Form(...),
    insurer_name: str = Form(...),
    plan_name: str = Form(None),
    sum_assured: str = Form(None),
    premium_amount: str = Form(None),
    premium_frequency: str = Form(None),
    policy_start_date: str = Form(...),
    policy_end_date: str = Form(None),
    maturity_amount: str = Form(None),
    status: str = Form("ACTIVE"),
    agent_name: str = Form(None),
    file_number: str = Form(None),
    insured_person_id: int = Form(None),
    holder_person_id: int = Form(None),
    fund_name: str = Form(None),
    fund_type: str = Form(None),
    premium_term: int = Form(None),
    payment_method: str = Form("MANUAL"),
    lapse_after_missed_payments: int = Form(2),
    notes: str = Form(None),
    db: Session = Depends(get_db),
):
    premium_end_date_calc = None
    if policy_start_date and premium_term:
        premium_end_date_calc = date.fromisoformat(policy_start_date) + relativedelta(years=premium_term)

    policy = InsurancePolicy(
        insured_person_id=insured_person_id,
        holder_person_id=holder_person_id,
        policy_number=policy_number,
        insurance_type=InsuranceType(insurance_type),
        insurer_name=insurer_name,
        plan_name=plan_name,
        sum_assured=Decimal(sum_assured) if sum_assured else None,
        premium_amount=Decimal(premium_amount) if premium_amount else None,
        premium_frequency=premium_frequency,
        policy_start_date=date.fromisoformat(policy_start_date),
        policy_end_date=date.fromisoformat(policy_end_date) if policy_end_date else None,
        premium_end_date=premium_end_date_calc,
        maturity_amount=Decimal(maturity_amount) if maturity_amount else None,
        status=InsuranceStatus(status),
        agent_name=agent_name,
        file_number=file_number,
        fund_name=fund_name,
        fund_type=fund_type,
        premium_term=premium_term,
        payment_method=PaymentMethod(payment_method),
        lapse_after_missed_payments=lapse_after_missed_payments,
        notes=notes,
    )
    db.add(policy)
    db.flush()  # To get policy.id

    # Automatic Generation of Payments till today (respecting term limit)
    generate_auto_payments(db, policy)

    db.commit()
    return RedirectResponse("/insurance/", status_code=303)


# ---------- POLICY DETAIL ----------
@router.get("/{policy_id}", response_class=HTMLResponse)
def policy_detail(policy_id: int, request: Request, db: Session = Depends(get_db)):
    policy = db.query(InsurancePolicy).filter(InsurancePolicy.id == policy_id).first()
    insured_person = db.query(Person).filter(Person.id == policy.insured_person_id).first()
    holder_person = db.query(Person).filter(Person.id == policy.holder_person_id).first()
    payments = db.query(PremiumPayment).filter(PremiumPayment.policy_id == policy_id).order_by(PremiumPayment.payment_date.desc()).all()
    nominees = db.query(InsuranceNominee).filter(InsuranceNominee.policy_id == policy_id).all()
    valuations = db.query(InsuranceValuation).filter(InsuranceValuation.policy_id == policy_id).order_by(InsuranceValuation.valuation_date.desc()).all()
    
    # Ensure premium_end_date is calculated if missing
    if not policy.premium_end_date and policy.policy_start_date and policy.premium_term:
        policy.premium_end_date = policy.policy_start_date + relativedelta(years=policy.premium_term)
        db.commit()

    # Calculations
    total_premium_to_be_paid = (policy.premium_amount or 0) * (policy.premium_term or 0)
    premium_paid_till_date = db.query(func.sum(PremiumPayment.amount)).filter(
        PremiumPayment.policy_id == policy_id,
        PremiumPayment.status == PaymentStatus.PAID,
        (PremiumPayment.transaction_type == "PREMIUM") | (PremiumPayment.transaction_type == None)
    ).scalar() or 0
    
    # Count how many premium payments have been made
    paid_premium_count = db.query(func.count(PremiumPayment.id)).filter(
        PremiumPayment.policy_id == policy_id,
        PremiumPayment.status == PaymentStatus.PAID,
        (PremiumPayment.transaction_type == "PREMIUM") | (PremiumPayment.transaction_type == None)
    ).scalar() or 0
    
    # Calculate the total expected payments for the term based on frequency
    expected_payments = policy.premium_term or 0
    if policy.premium_frequency == "MONTHLY":
        expected_payments = (policy.premium_term or 0) * 12
    elif policy.premium_frequency == "QUARTERLY":
        expected_payments = (policy.premium_term or 0) * 4
    elif policy.premium_frequency == "HALF_YEARLY":
        expected_payments = (policy.premium_term or 0) * 2
    elif policy.premium_frequency == "SINGLE":
        expected_payments = 1

    # If all premiums are paid, treat total paid = target (differences are GST/riders)
    if expected_payments > 0 and paid_premium_count >= expected_payments:
        premium_paid_till_date = total_premium_to_be_paid
    
    balance_to_be_paid = max(0, total_premium_to_be_paid - premium_paid_till_date)

    # Missed payments calculation
    missed_payments_count = db.query(func.count(PremiumPayment.id)).filter(
        PremiumPayment.policy_id == policy_id,
        PremiumPayment.status == PaymentStatus.UNPAID,
        PremiumPayment.payment_date <= date.today()
    ).scalar() or 0

    # Term calculations
    years_completed = 0
    remaining_term = 0
    if policy.policy_start_date:
        today = date.today()
        # Years completed is capped at premium_term
        diff = relativedelta(today, policy.policy_start_date)
        years_completed = min(policy.premium_term or 0, diff.years + (diff.months / 12.0))
        remaining_term = max(0, (policy.premium_term or 0) - years_completed)

    # Calculate Policy XIRR & Latest Surrender Value
    policy_xirr = None
    latest_surrender_value = None
    
    paid_payments = [p for p in payments if p.status.value == "PAID"]
    latest_val = valuations[0] if valuations else None
    
    if latest_val:
        latest_surrender_value = latest_val.surrender_value or latest_val.total_value or latest_val.fund_value
        
    if policy.status.value == "SURRENDERED" and policy.surrendered_amount is not None:
        latest_surrender_value = policy.surrendered_amount

    # XIRR: build cash flows from all paid transactions
    # PREMIUM = outflow (negative), PARTIAL_WITHDRAWAL/SURRENDER/MATURITY_PAYOUT = inflow (positive)
    inflow_types = {'PARTIAL_WITHDRAWAL', 'SURRENDER', 'MATURITY_PAYOUT'}
    if paid_payments:
        dates = []
        flows = []
        has_terminal_payout = False
        for p in paid_payments:
            txn_type = getattr(p, 'transaction_type', None) or 'PREMIUM'
            if txn_type in {'PARTIAL_WITHDRAWAL', 'SURRENDER', 'MATURITY_PAYOUT'}:
                flows.append(float(p.amount))   # inflow: positive
                if txn_type in {'SURRENDER', 'MATURITY_PAYOUT'}:
                    has_terminal_payout = True
            else:
                flows.append(-float(p.amount))  # outflow: negative
            dates.append(p.payment_date)
        
        # If no terminal payout recorded, use latest valuation as terminal value
        if not has_terminal_payout:
            inflow = None
            inflow_date = None
            if policy.status.value == "SURRENDERED" and policy.surrendered_amount is not None:
                inflow = float(policy.surrendered_amount)
                inflow_date = latest_val.valuation_date if latest_val else date.today()
            elif policy.status.value == "MATURED" and policy.maturity_amount is not None:
                inflow = float(policy.maturity_amount)
                inflow_date = policy.policy_end_date if policy.policy_end_date else (latest_val.valuation_date if latest_val else date.today())
            elif latest_surrender_value is not None:
                inflow = float(latest_surrender_value)
                inflow_date = latest_val.valuation_date if latest_val else date.today()
            if inflow is not None and inflow > 0:
                dates.append(inflow_date)
                flows.append(inflow)
            
        from services.performance_engine import calculate_xirr
        try:
            xirr_val = calculate_xirr(dates, flows)
            if xirr_val is not None:
                policy_xirr = xirr_val * 100
        except Exception as e:
            print(f"Error calculating policy XIRR: {e}")

    # Lock Status calculation
    lock_status = "—"
    if policy.policy_start_date and policy.premium_term:
        lock_release_date = policy.policy_start_date + relativedelta(years=policy.premium_term + 1)
        if date.today() >= lock_release_date:
            lock_status = "UNLOCKED"
        else:
            lock_status = "LOCKED"

    # Available Withdrawal calculation
    # For policies which are locked, keep withdrawable as blank (None)
    withdrawable = None
    if lock_status != "LOCKED":
        fund_val_for_withdrawable = None
        if latest_val:
            fund_val_for_withdrawable = latest_val.fund_value or latest_val.total_value or latest_val.surrender_value
        if fund_val_for_withdrawable is not None and policy.premium_amount is not None:
            calc_val = float(fund_val_for_withdrawable) - 1.5 * float(policy.premium_amount)
            withdrawable = max(0.0, calc_val)

    return templates.TemplateResponse(
        request=request,
        name="insurance/detail.html",
        context={
            "policy": policy,
            "insured_person": insured_person,
            "holder_person": holder_person,
            "payments": payments,
            "nominees": nominees,
            "valuations": valuations,
            "total_premium_to_be_paid": total_premium_to_be_paid,
            "premium_paid_till_date": premium_paid_till_date,
            "balance_to_be_paid": balance_to_be_paid,
            "years_completed": round(years_completed, 1),
            "remaining_term": round(remaining_term, 1),
            "missed_payments_count": missed_payments_count,
            "payment_statuses": [e.value for e in PaymentStatus],
            "withdrawable": withdrawable,
            "policy_xirr": policy_xirr,
            "lock_status": lock_status,
            "latest_surrender_value": latest_surrender_value,
        }
    )


# ---------- ADD PREMIUM PAYMENT ----------
@router.post("/{policy_id}/add_payment")
def add_payment(
    policy_id: int,
    transaction_type: str = Form("PREMIUM"),
    payment_date: str = Form(...),
    amount: str = Form(...),
    payment_mode: str = Form(None),
    receipt_number: str = Form(None),
    notes: str = Form(None),
    db: Session = Depends(get_db),
):
    payment = PremiumPayment(
        policy_id=policy_id,
        transaction_type=transaction_type,
        payment_date=date.fromisoformat(payment_date),
        amount=Decimal(amount),
        payment_mode=payment_mode,
        receipt_number=receipt_number,
        notes=notes,
    )
    db.add(payment)
    db.commit()
    return RedirectResponse(f"/insurance/{policy_id}", status_code=303)


@router.post("/payment/{payment_id}/update")
def update_payment_details(
    payment_id: int,
    transaction_type: str = Form("PREMIUM"),
    payment_date: str = Form(...),
    amount: str = Form(...),
    payment_mode: str = Form(None),
    status: str = Form(...),
    receipt_number: str = Form(None),
    notes: str = Form(None),
    db: Session = Depends(get_db)
):
    try:
        payment = db.query(PremiumPayment).filter(PremiumPayment.id == payment_id).first()
        if not payment:
            return {"success": False, "error": "Payment record not found."}
            
        payment.transaction_type = transaction_type or "PREMIUM"
        payment.payment_date = date.fromisoformat(payment_date)
        
        # Clean amount string
        clean_amount = amount.replace("₹", "").replace(",", "").strip()
        payment.amount = Decimal(clean_amount)
        
        payment.payment_mode = payment_mode or None
        payment.status = PaymentStatus(status)
        payment.receipt_number = receipt_number or None
        payment.notes = notes or None
        
        db.commit()
        return {"success": True, "message": "Payment updated successfully."}
    except Exception as e:
        db.rollback()
        return {"success": False, "error": str(e)}

@router.post("/payment/{payment_id}/update_status")
def update_payment_status(
    payment_id: int,
    status: str = Form(...),
    db: Session = Depends(get_db),
):
    payment = db.query(PremiumPayment).filter(PremiumPayment.id == payment_id).first()
    if payment:
        payment.status = PaymentStatus(status)
        db.commit()
        return RedirectResponse(f"/insurance/{payment.policy_id}", status_code=303)
    return RedirectResponse("/insurance/", status_code=303)
@router.post("/payment/{payment_id}/delete")
def delete_payment(
    payment_id: int,
    db: Session = Depends(get_db),
):
    payment = db.query(PremiumPayment).filter(PremiumPayment.id == payment_id).first()
    if payment:
        policy_id = payment.policy_id
        db.delete(payment)
        db.commit()
        return RedirectResponse(f"/insurance/{policy_id}", status_code=303)
    return RedirectResponse("/insurance/", status_code=303)


# ---------- ADD NOMINEE ----------
@router.post("/{policy_id}/add_nominee")
def add_nominee(
    policy_id: int,
    nominee_name: str = Form(...),
    relationship: str = Form(None),
    share_percentage: str = Form(None),
    is_minor: bool = Form(False),
    appointee_name: str = Form(None),
    db: Session = Depends(get_db),
):
    nominee = InsuranceNominee(
        policy_id=policy_id,
        nominee_name=nominee_name,
        relationship=relationship,
        share_percentage=Decimal(share_percentage) if share_percentage else None,
        is_minor=is_minor,
        appointee_name=appointee_name,
    )
    db.add(nominee)
    db.commit()
    return RedirectResponse(f"/insurance/{policy_id}", status_code=303)


# ---------- ADD VALUATION ----------
@router.post("/{policy_id}/add_valuation")
def add_valuation(
    policy_id: int,
    valuation_date: str = Form(...),
    surrender_value: str = Form(None),
    fund_value: str = Form(None),
    bonus_accrued: str = Form(None),
    total_value: str = Form(None),
    notes: str = Form(None),
    db: Session = Depends(get_db),
):
    valuation = InsuranceValuation(
        policy_id=policy_id,
        valuation_date=date.fromisoformat(valuation_date),
        surrender_value=Decimal(surrender_value) if surrender_value else None,
        fund_value=Decimal(fund_value) if fund_value else None,
        bonus_accrued=Decimal(bonus_accrued) if bonus_accrued else None,
        total_value=Decimal(total_value) if total_value else None,
        notes=notes,
    )
    db.add(valuation)
    db.commit()
    return RedirectResponse(f"/insurance/{policy_id}", status_code=303)


# ---------- UPDATE POLICY ----------
@router.post("/{policy_id}/update")
def update_policy(
    policy_id: int,
    file_number: str = Form(None),
    policy_number: str = Form(None),
    plan_name: str = Form(None),
    fund_name: str = Form(None),
    fund_type: str = Form(None),
    sum_assured: str = Form(None),
    premium_amount: str = Form(None),
    status: str = Form(None),
    premium_term: str = Form(None),
    premium_frequency: str = Form(None),
    payment_method: str = Form(None),
    policy_start_date: str = Form(None),
    policy_end_date: str = Form(None),
    premium_end_date: str = Form(None),
    maturity_amount: str = Form(None),
    surrendered_amount: str = Form(None),
    agent_name: str = Form(None),
    notes: str = Form(None),
    db: Session = Depends(get_db)
):
    try:
        policy = db.query(InsurancePolicy).filter(InsurancePolicy.id == policy_id).first()
        if not policy:
            return {"success": False, "error": "Policy not found."}
        
        # Capture old values
        old_freq = policy.premium_frequency
        old_amount = policy.premium_amount
        old_start_date = policy.policy_start_date
        old_premium_end_date = policy.premium_end_date
        old_payment_method = policy.payment_method
        old_term = policy.premium_term

        if file_number is not None:
            policy.file_number = file_number or None
        if policy_number is not None:
            policy.policy_number = policy_number
        if plan_name is not None:
            policy.plan_name = plan_name or None
        if fund_name is not None:
            policy.fund_name = fund_name or None
        if fund_type is not None:
            policy.fund_type = fund_type or None
        if sum_assured is not None:
            policy.sum_assured = Decimal(sum_assured.replace("₹", "").replace(",", "").strip()) if sum_assured else None
        if premium_amount is not None:
            policy.premium_amount = Decimal(premium_amount.replace("₹", "").replace(",", "").strip()) if premium_amount else None
        if status is not None:
            policy.status = InsuranceStatus(status)
        if premium_term is not None:
            policy.premium_term = int(premium_term) if premium_term else None
        if premium_frequency is not None:
            policy.premium_frequency = premium_frequency or None
        if payment_method is not None:
            policy.payment_method = PaymentMethod(payment_method)
        if policy_start_date is not None:
            policy.policy_start_date = date.fromisoformat(policy_start_date) if policy_start_date else None
        if policy_end_date is not None:
            policy.policy_end_date = date.fromisoformat(policy_end_date) if policy_end_date else None
        if premium_end_date is not None:
            policy.premium_end_date = date.fromisoformat(premium_end_date) if premium_end_date else None
        if maturity_amount is not None:
            policy.maturity_amount = Decimal(maturity_amount.replace("₹", "").replace(",", "").strip()) if maturity_amount else None
        if surrendered_amount is not None:
            policy.surrendered_amount = Decimal(surrendered_amount.replace("₹", "").replace(",", "").strip()) if surrendered_amount else None
        if agent_name is not None:
            policy.agent_name = agent_name or None
        if notes is not None:
            policy.notes = notes or None
            
        # Recalculate premium_end_date if start date or term changed
        if (policy.policy_start_date != old_start_date) or (policy.premium_term != old_term):
            if policy.policy_start_date and policy.premium_term:
                policy.premium_end_date = policy.policy_start_date + relativedelta(years=policy.premium_term)

        # Check if fields affecting auto payments have changed
        freq_changed = (policy.premium_frequency != old_freq)
        amount_changed = (policy.premium_amount != old_amount)
        start_date_changed = (policy.policy_start_date != old_start_date)
        end_date_changed = (policy.premium_end_date != old_premium_end_date)
        method_changed = (policy.payment_method != old_payment_method)

        if freq_changed or amount_changed or start_date_changed or end_date_changed or method_changed:
            # Delete old auto-generated payments
            db.query(PremiumPayment).filter(
                PremiumPayment.policy_id == policy.id,
                PremiumPayment.notes == "Automatically generated on policy creation"
            ).delete()
            
            # Generate new auto-generated payments
            generate_auto_payments(db, policy)

        db.commit()
        return {"success": True, "message": "Policy details updated successfully."}
    except Exception as e:
        db.rollback()
        return {"success": False, "error": str(e)}


# ---------- UPDATE VALUATION ----------
@router.post("/valuation/{valuation_id}/update")
def update_valuation(
    valuation_id: int,
    valuation_date: str = Form(...),
    surrender_value: str = Form(None),
    fund_value: str = Form(None),
    bonus_accrued: str = Form(None),
    total_value: str = Form(None),
    notes: str = Form(None),
    db: Session = Depends(get_db),
):
    try:
        val = db.query(InsuranceValuation).filter(InsuranceValuation.id == valuation_id).first()
        if not val:
            return {"success": False, "error": "Valuation record not found."}
            
        val.valuation_date = date.fromisoformat(valuation_date)
        
        def clean_decimal(v):
            if not v or v.strip() == "" or v.strip() == "—":
                return None
            return Decimal(v.replace("₹", "").replace(",", "").strip())

        val.surrender_value = clean_decimal(surrender_value)
        val.fund_value = clean_decimal(fund_value)
        val.bonus_accrued = clean_decimal(bonus_accrued)
        val.total_value = clean_decimal(total_value)
        val.notes = notes or None
        
        db.commit()
        return {"success": True, "message": "Valuation updated successfully."}
    except Exception as e:
        db.rollback()
        return {"success": False, "error": str(e)}


# ---------- DELETE VALUATION ----------
@router.post("/valuation/{valuation_id}/delete")
def delete_valuation(
    valuation_id: int,
    db: Session = Depends(get_db),
):
    val = db.query(InsuranceValuation).filter(InsuranceValuation.id == valuation_id).first()
    if val:
        policy_id = val.policy_id
        db.delete(val)
        db.commit()
        return RedirectResponse(f"/insurance/{policy_id}", status_code=303)
    return RedirectResponse("/insurance/", status_code=303)