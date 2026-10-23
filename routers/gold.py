# routers/gold.py
import os
import shutil
import tempfile
import re
import datetime
import io
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from fastapi import APIRouter, Depends, Request, Form, File, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy.orm import Session
from database import get_db
from models.gold import GoldHolding
from models.persons import Person
from models.accounts import Account
from services.gold_importer import import_sgb_from_excel
from decimal import Decimal

router = APIRouter(prefix="/gold", tags=["gold"])
templates = Jinja2Templates(directory="templates")

# Register format_inr filter
from routers.mutual_funds import format_inr
templates.env.filters["format_inr"] = format_inr

def format_sgb_k(val):
    try:
        f_val = float(val)
        return f"₹{f_val / 1000:.1f} K"
    except (ValueError, TypeError):
        return "—"

def format_sgb_l(val):
    try:
        f_val = float(val)
        sign = "-" if f_val < 0 else ""
        return f"{sign}₹{abs(f_val) / 100000:.2f} L"
    except (ValueError, TypeError):
        return "—"

templates.env.filters["format_sgb_k"] = format_sgb_k
templates.env.filters["format_sgb_l"] = format_sgb_l

def get_broker_short_code(broker_name: str) -> str:
    name = broker_name.upper()
    if "NUVAMA" in name:
        return "NUV"
    elif "JM" in name:
        return "JM"
    elif "CENTRUM" in name:
        return "CENT"
    elif "PHYSICAL" in name:
        return "Physical"
    return broker_name

def sync_latest_sgb_from_global_report(db: Session):
    """Auto-detects and imports the latest Global Positions Report if not yet imported or if newer."""
    try:
        from services.consolidated_portfolio_engine import get_latest_global_positions_file
        from services.gold_importer import import_sgb_from_excel, parse_report_date
        from models.gold import GoldHolding, GoldType
        
        latest_file = get_latest_global_positions_file()
        if not latest_file:
            return
            
        file_date = parse_report_date(latest_file)
        if not file_date:
            return
            
        count_for_date = db.query(GoldHolding).filter(
            GoldHolding.gold_type == GoldType.SGB,
            GoldHolding.as_of_date == file_date
        ).count()
        
        if count_for_date == 0:
            import_sgb_from_excel(db, latest_file)
    except Exception as e:
        print(f"Note: SGB auto-sync from Global Positions Report encountered: {e}")

def get_flat_holdings_list(db: Session) -> list:
    import re
    import glob
    from datetime import datetime, date, timedelta
    from sqlalchemy import func
    from models.gold import GoldHolding, GoldType
    from models.persons import Person
    from models.accounts import Account
    
    # Auto-sync latest Global Positions Report if needed
    sync_latest_sgb_from_global_report(db)
    
    latest_date = db.query(func.max(GoldHolding.as_of_date)).filter(GoldHolding.gold_type == GoldType.SGB).scalar()
    
    if latest_date:
        db_holdings = db.query(GoldHolding).filter(GoldHolding.gold_type == GoldType.SGB, GoldHolding.as_of_date == latest_date).all()
    else:
        db_holdings = db.query(GoldHolding).filter(GoldHolding.gold_type == GoldType.SGB).all()
        
    persons_map = {p.id: p.short_name for p in db.query(Person).all()}
    accounts_map = {a.id: a for a in db.query(Account).all()}
    
    fallback_gold_price = max([h.last_price for h in db_holdings if h.last_price is not None] or [Decimal("7000")])
    
    # Build active holdings lookup from DB (which reflects the latest Global Positions Report)
    active_holdings = {}
    isin_lookup = {}
    
    for h in db_holdings:
        p_name = persons_map.get(h.person_id)
        acct = accounts_map.get(h.account_id)
        b_name = acct.institution_name if acct else ""
        if "NUVAMA" in b_name.upper():
            b_code = "NUV"
        elif "JM" in b_name.upper():
            b_code = "JM"
        elif "CENTRUM" in b_name.upper():
            b_code = "CENT"
        else:
            b_code = b_name
            
        key = (p_name, b_code, h.isin)
        active_holdings[key] = {
            "holding": h,
            "owner": p_name,
            "broker": b_code,
            "isin": h.isin,
            "series": h.sgb_series,
            "description": h.description,
            "maturity_date": h.sgb_maturity_date,
            "units": float(h.sgb_units or 0),
            "total_cost": float(h.total_cost or 0),
            "purchase_price": float(h.purchase_price_per_gram or (h.total_cost / h.sgb_units if h.sgb_units else 0)),
            "last_price": float(h.last_price or fallback_gold_price),
            "purchase_date": h.purchase_date
        }
        
        isin_lookup[h.isin] = {
            "series": h.sgb_series,
            "maturity_date": h.sgb_maturity_date,
            "last_price": h.last_price
        }
    
    def parse_sgb_name(name_str: str):
        # SGB ISINs start with IN00 (Government of India securities)
        isin_match = re.search(r'\b(IN[A-Z0-9]{10})\b', name_str)
        isin = isin_match.group(1) if isin_match else None
        
        series_match = re.search(r'\b(SR-[IVXLCDM]+\s+\d{4}-\d{2})\b', name_str, re.IGNORECASE)
        series = series_match.group(1) if series_match else None
        
        maturity_date = None
        month_year_match = re.search(r'\b(JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|OCT|NOV|DEC|JUNE)\s+(\d{4})\b', name_str.upper())
        if month_year_match:
            month_str = month_year_match.group(1)
            year_str = month_year_match.group(2)
            if month_str == "JUNE":
                month_str = "JUN"
            try:
                dt = datetime.strptime(f"01-{month_str}-{year_str}", "%d-%b-%Y")
                maturity_date = dt.date()
            except ValueError:
                pass
                
        return isin, series, maturity_date
        
    def is_primary_purchase(purchase_date, maturity_date) -> bool:
        if not purchase_date or not maturity_date:
            return False
        issuance_year = maturity_date.year - 8
        return purchase_date.year == issuance_year and purchase_date.month == maturity_date.month
        
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    folder = os.path.join(base_dir, "DATA", "SGB GOLD BOND DATA")
    files = glob.glob(os.path.join(folder, "*.xlsx"))
    
    # Store parsed active broker lots: (owner, broker, isin) -> list of lot dicts
    broker_lots = {}
    
    for f in files:
        filename = os.path.basename(f)
        parts = filename.split('_')
        owner = parts[0]
        broker_code = parts[1]
        
        if broker_code == "CBL":
            broker_code = "CENT"
        elif broker_code == "NUVAMA":
            broker_code = "NUV"
            
        try:
            wb = openpyxl.load_workbook(f, data_only=True)
            ws = wb.active
        except Exception:
            continue
            
        current_bond_name = None
        current_isin = None
        current_series = None
        current_maturity = None
        current_txns = []
        closing_balance_qty = None
        
        def process_block(b_owner, b_broker, b_isin, b_series, b_mat, b_name, txns, closing_qty):
            if not b_isin or not txns:
                return
            key = (b_owner, b_broker, b_isin)
            
            # Cross-check: strictly skip if not in active DB / Global Positions Report holdings
            if key not in active_holdings:
                return
                
            # If closing balance is explicitly 0, skip sold bond
            if closing_qty is not None and closing_qty <= 0:
                return
                
            # Process buy and sell transactions (FIFO lot matching)
            buy_lots = []
            for t in txns:
                if t["type"] == "Buy" and t["qty"] > 0:
                    buy_lots.append({
                        "date": t["date"],
                        "qty": t["qty"],
                        "cost": t["cost"],
                        "price_per_unit": t["cost"] / t["qty"] if t["qty"] > 0 else 0
                    })
                elif t["type"] == "Sell" and t["qty"] > 0:
                    sell_rem = t["qty"]
                    while sell_rem > 0 and buy_lots:
                        if buy_lots[0]["qty"] <= sell_rem:
                            sell_rem -= buy_lots[0]["qty"]
                            buy_lots.pop(0)
                        else:
                            buy_lots[0]["qty"] -= sell_rem
                            buy_lots[0]["cost"] = buy_lots[0]["qty"] * buy_lots[0]["price_per_unit"]
                            sell_rem = 0
                            
            if buy_lots:
                if key not in broker_lots:
                    broker_lots[key] = []
                for lot in buy_lots:
                    broker_lots[key].append({
                        "owner": b_owner,
                        "broker": b_broker,
                        "isin": b_isin,
                        "series": b_series,
                        "description": b_name,
                        "maturity_date": b_mat,
                        "purchase_date": lot["date"],
                        "units": lot["qty"],
                        "total_cost": lot["cost"]
                    })
                    
        for r in range(1, ws.max_row+1):
            cell_c = ws.cell(row=r, column=3).value
            if cell_c and str(cell_c).startswith("Name: "):
                process_block(owner, broker_code, current_isin, current_series, current_maturity, current_bond_name, current_txns, closing_balance_qty)
                
                current_bond_name = str(cell_c).replace("Name: ", "").strip()
                current_isin, current_series, current_maturity = parse_sgb_name(current_bond_name)
                current_txns = []
                closing_balance_qty = None
                
                if current_isin in isin_lookup:
                    db_info = isin_lookup[current_isin]
                    if not current_series:
                        current_series = db_info["series"]
                    if not current_maturity:
                        current_maturity = db_info["maturity_date"]
                        
            cell_e = ws.cell(row=r, column=5).value
            if cell_e in ["Buy", "Sell"] and current_isin:
                dt_val = ws.cell(row=r, column=3).value
                p_dt = None
                if isinstance(dt_val, datetime):
                    p_dt = dt_val.date()
                elif isinstance(dt_val, date):
                    p_dt = dt_val
                qty_val = float(ws.cell(row=r, column=7).value or 0)
                cost_val = float(ws.cell(row=r, column=11).value or 0)
                current_txns.append({
                    "type": cell_e,
                    "date": p_dt,
                    "qty": qty_val,
                    "cost": cost_val
                })
            elif cell_c == "Closing Balance":
                closing_balance_qty = float(ws.cell(row=r, column=7).value or 0)
                
        # Process final block in file
        process_block(owner, broker_code, current_isin, current_series, current_maturity, current_bond_name, current_txns, closing_balance_qty)
    
    # Assemble final flat holdings list cross-checked against active holdings
    holdings_list = []
    idx = 1
    
    for key in sorted(active_holdings.keys(), key=lambda k: (k[0], k[1], active_holdings[k]["maturity_date"] or date.max)):
        active_info = active_holdings[key]
        lots = broker_lots.get(key)
        
        if lots and sum(l["units"] for l in lots) > 0:
            for lot in lots:
                units = lot["units"]
                total_cost = lot["total_cost"]
                purchase_date_dt = lot["purchase_date"]
                maturity_date_dt = active_info["maturity_date"]
                
                purchase_date_str = purchase_date_dt.strftime("%d-%b-%Y") if purchase_date_dt else "—"
                maturity_date_str = maturity_date_dt.strftime("%d-%b-%Y") if maturity_date_dt else "—"
                
                is_long_term = "—"
                days_to_lt = 0
                if purchase_date_dt and latest_date:
                    try:
                        date_2_years = purchase_date_dt.replace(year=purchase_date_dt.year + 2)
                    except ValueError:
                        date_2_years = purchase_date_dt + timedelta(days=2*365)
                    is_long_term = "Yes" if latest_date >= date_2_years else "No"
                    delta = (date_2_years - latest_date).days
                    if delta > 0:
                        days_to_lt = delta
                        
                current_price = active_info["last_price"]
                purchase_price = total_cost / units if units > 0 else 0.0
                current_value = units * current_price
                pnl = current_value - total_cost
                return_pct = pnl / total_cost if total_cost > 0 else 0.0
                
                if is_primary_purchase(purchase_date_dt, maturity_date_dt):
                    tax_label = "0%"
                else:
                    tax_label = "12.5%" if is_long_term == "Yes" else "35%"
                    
                holdings_list.append({
                    "id": idx,
                    "owner": active_info["owner"],
                    "broker": active_info["broker"],
                    "tax": tax_label,
                    "description": active_info["description"],
                    "isin": active_info["isin"],
                    "series": active_info["series"] or active_info["isin"] or "—",
                    "maturity_date": maturity_date_str,
                    "units": units,
                    "purchase_price": purchase_price,
                    "total_cost": total_cost,
                    "current_price": current_price,
                    "current_value": current_value,
                    "pnl": pnl,
                    "return_pct": return_pct,
                    "purchase_date": purchase_date_str,
                    "is_long_term": is_long_term,
                    "days_to_lt": days_to_lt
                })
                idx += 1
        else:
            # Direct DB holding (e.g. Primary allotment without individual broker trade logs)
            units = active_info["units"]
            total_cost = active_info["total_cost"]
            purchase_date_dt = active_info["purchase_date"]
            maturity_date_dt = active_info["maturity_date"]
            
            purchase_date_str = purchase_date_dt.strftime("%d-%b-%Y") if purchase_date_dt else "—"
            maturity_date_str = maturity_date_dt.strftime("%d-%b-%Y") if maturity_date_dt else "—"
            
            is_long_term = "—"
            days_to_lt = 0
            if purchase_date_dt and latest_date:
                try:
                    date_2_years = purchase_date_dt.replace(year=purchase_date_dt.year + 2)
                except ValueError:
                    date_2_years = purchase_date_dt + timedelta(days=2*365)
                is_long_term = "Yes" if latest_date >= date_2_years else "No"
                delta = (date_2_years - latest_date).days
                if delta > 0:
                    days_to_lt = delta
                    
            current_price = active_info["last_price"]
            purchase_price = active_info["purchase_price"]
            current_value = units * current_price
            pnl = current_value - total_cost
            return_pct = pnl / total_cost if total_cost > 0 else 0.0
            
            if is_primary_purchase(purchase_date_dt, maturity_date_dt) or active_info["broker"] == "NUV":
                tax_label = "0%"
            else:
                tax_label = "12.5%" if is_long_term == "Yes" else "35%"
                
            holdings_list.append({
                "id": idx,
                "owner": active_info["owner"],
                "broker": active_info["broker"],
                "tax": tax_label,
                "description": active_info["description"],
                "isin": active_info["isin"],
                "series": active_info["series"] or active_info["isin"] or "—",
                "maturity_date": maturity_date_str,
                "units": units,
                "purchase_price": purchase_price,
                "total_cost": total_cost,
                "current_price": current_price,
                "current_value": current_value,
                "pnl": pnl,
                "return_pct": return_pct,
                "purchase_date": purchase_date_str,
                "is_long_term": is_long_term,
                "days_to_lt": days_to_lt
            })
            idx += 1
            
    return holdings_list

from sqlalchemy import text

def run_gold_migrations(db: Session):
    migrations = [
        "ALTER TABLE gold.holdings ADD COLUMN IF NOT EXISTS account_id INTEGER REFERENCES accounts.accounts(id)",
        "ALTER TABLE gold.holdings ADD COLUMN IF NOT EXISTS isin VARCHAR(20)",
        "ALTER TABLE gold.holdings ADD COLUMN IF NOT EXISTS last_price NUMERIC(18,2)",
        "ALTER TABLE gold.holdings ADD COLUMN IF NOT EXISTS current_value NUMERIC(18,2)",
        "ALTER TABLE gold.holdings ADD COLUMN IF NOT EXISTS as_of_date DATE",
    ]
    for sql in migrations:
        try:
            db.execute(text(sql))
        except Exception as e:
            print(f"Migration note: {e}")
    db.commit()

PERSONS = ["PATRIARCH", "MEMBER_1", "MEMBER_2", "MEMBER_3"]

def load_primary_sgb_map() -> dict:
    excel_path = "SGB LIST.xlsx"
    primary_data = {}
    
    if os.path.exists(excel_path):
        try:
            df = pd.read_excel(excel_path, sheet_name="Sheet3", header=None)
            for r_idx in range(5, 27):
                row = df.iloc[r_idx]
                mat_date = row[11]
                if pd.isna(mat_date):
                    continue
                if isinstance(mat_date, str):
                    mat_date = pd.to_datetime(mat_date)
                y = mat_date.year
                if y < 2000:
                    mat_date = mat_date.replace(year=y + 100)
                mat_key = mat_date.strftime("%Y-%m")
                
                cols = {
                    "PATRIARCH": (12, 13),
                    "MEMBER_1": (14, 15),
                    "MEMBER_2": (16, 17),
                    "MEMBER_3": (18, 19)
                }
                
                primary_data[mat_key] = {}
                for p_name, (q_col, a_col) in cols.items():
                    qty = row[q_col]
                    avg = row[a_col]
                    if pd.notna(qty) and float(qty) > 0:
                        primary_data[mat_key][p_name] = {
                            "qty": Decimal(str(qty)),
                            "avg": Decimal(str(avg))
                        }
            return primary_data
        except Exception as e:
            print(f"Error parsing SGB LIST.xlsx Sheet3: {e}. Falling back to default map.")
            
    primary_data = {
        "2028-07": {
            "MEMBER_1": {"qty": Decimal("1000"), "avg": Decimal("4852")},
            "MEMBER_2": {"qty": Decimal("1000"), "avg": Decimal("4681")}
        },
        "2028-08": {
            "PATRIARCH": {"qty": Decimal("946"), "avg": Decimal("5284")},
            "MEMBER_1": {"qty": Decimal("946"), "avg": Decimal("5334")}
        },
        "2029-03": {
            "MEMBER_1": {"qty": Decimal("1000"), "avg": Decimal("4612")},
            "MEMBER_2": {"qty": Decimal("1000"), "avg": Decimal("4621")}
        },
        "2029-06": {
            "MEMBER_2": {"qty": Decimal("2000"), "avg": Decimal("4792")},
            "MEMBER_3": {"qty": Decimal("2000"), "avg": Decimal("4792")}
        },
        "2031-12": {
            "MEMBER_1": {"qty": Decimal("1500"), "avg": Decimal("6149")},
            "MEMBER_2": {"qty": Decimal("1000"), "avg": Decimal("6149")}
        },
        "2032-02": {
            "PATRIARCH": {"qty": Decimal("1000"), "avg": Decimal("6263")},
            "MEMBER_1": {"qty": Decimal("1000"), "avg": Decimal("6213")}
        }
    }
    return primary_data

def compile_sgb_tables(db: Session) -> dict:
    from models.enums import GoldType
    from sqlalchemy import func
    
    # Auto-sync latest Global Positions Report if needed
    sync_latest_sgb_from_global_report(db)
    
    # Run migrations just in case
    run_gold_migrations(db)
    
    latest_date = db.query(func.max(GoldHolding.as_of_date)).filter(GoldHolding.gold_type == GoldType.SGB).scalar()
    
    if latest_date:
        holdings = db.query(GoldHolding).filter(GoldHolding.gold_type == GoldType.SGB, GoldHolding.as_of_date == latest_date).all()
    else:
        holdings = db.query(GoldHolding).filter(GoldHolding.gold_type == GoldType.SGB).all()
        
    persons_map = {p.id: p.short_name for p in db.query(Person).all()}
    accounts_map = {a.id: a for a in db.query(Account).all()}
    prim_map = load_primary_sgb_map()
    
    issues_consolidated = {}
    issues_primary = {}
    issues_secondary = {}
    
    brokers_list = ["JM Financial", "Nuvama Wealth", "Centrum Finverse"]
    broker_primary = {b: {p: {"qty": Decimal("0"), "cost": Decimal("0")} for p in PERSONS} for b in brokers_list}
    broker_secondary = {b: {p: {"qty": Decimal("0"), "cost": Decimal("0")} for p in PERSONS} for b in brokers_list}
    
    for h in holdings:
        p_name = persons_map.get(h.person_id)
        if p_name not in PERSONS:
            continue
            
        key = (h.isin, h.sgb_series, h.sgb_maturity_date)
        mat_key = h.sgb_maturity_date.strftime("%Y-%m") if h.sgb_maturity_date else ""
        account = accounts_map.get(h.account_id)
        broker = account.institution_name if account else "Physical"
        
        if "JM" in broker.upper():
            broker_key = "JM Financial"
        elif "NUVAMA" in broker.upper():
            broker_key = "Nuvama Wealth"
        elif "CENTRUM" in broker.upper():
            broker_key = "Centrum Finverse"
        else:
            broker_key = broker
            if broker_key not in broker_primary:
                broker_primary[broker_key] = {p: {"qty": Decimal("0"), "cost": Decimal("0")} for p in PERSONS}
                broker_secondary[broker_key] = {p: {"qty": Decimal("0"), "cost": Decimal("0")} for p in PERSONS}
        
        qty = h.sgb_units or Decimal("0")
        total_cost = h.total_cost or Decimal("0")
        avg_price = h.purchase_price_per_gram or Decimal("0")
        last_price = h.last_price or Decimal("0")
        
        # Consolidated
        if key not in issues_consolidated:
            issues_consolidated[key] = {
                "series": h.sgb_series,
                "maturity_date": h.sgb_maturity_date,
                "isin": h.isin,
                "last_price": last_price,
                "holdings": {p: {"qty": Decimal("0"), "cost": Decimal("0")} for p in PERSONS}
            }
        issues_consolidated[key]["holdings"][p_name]["qty"] += qty
        issues_consolidated[key]["holdings"][p_name]["cost"] += total_cost
        
        # Primary Split logic
        prim_record = prim_map.get(mat_key, {}).get(p_name)
        is_primary = False
        prim_qty = Decimal("0")
        prim_avg = Decimal("0")
        
        if broker_key == "Nuvama Wealth":
            is_primary = True
            prim_qty = qty
            prim_avg = avg_price
        elif broker_key == "JM Financial" and prim_record:
            is_primary = True
            prim_qty = prim_record["qty"]
            prim_avg = prim_record["avg"]
            if prim_qty > qty:
                prim_qty = qty
                
        if is_primary and prim_qty > 0:
            prim_cost = prim_qty * prim_avg
            if key not in issues_primary:
                issues_primary[key] = {
                    "series": h.sgb_series,
                    "maturity_date": h.sgb_maturity_date,
                    "isin": h.isin,
                    "last_price": last_price,
                    "holdings": {p: {"qty": Decimal("0"), "cost": Decimal("0")} for p in PERSONS}
                }
            issues_primary[key]["holdings"][p_name]["qty"] += prim_qty
            issues_primary[key]["holdings"][p_name]["cost"] += prim_cost
            broker_primary[broker_key][p_name]["qty"] += prim_qty
            broker_primary[broker_key][p_name]["cost"] += prim_cost
            
            sec_qty = qty - prim_qty
            if sec_qty > 0:
                sec_cost = total_cost - prim_cost
                if key not in issues_secondary:
                    issues_secondary[key] = {
                        "series": h.sgb_series,
                        "maturity_date": h.sgb_maturity_date,
                        "isin": h.isin,
                        "last_price": last_price,
                        "holdings": {p: {"qty": Decimal("0"), "cost": Decimal("0")} for p in PERSONS}
                    }
                issues_secondary[key]["holdings"][p_name]["qty"] += sec_qty
                issues_secondary[key]["holdings"][p_name]["cost"] += sec_cost
                broker_secondary[broker_key][p_name]["qty"] += sec_qty
                broker_secondary[broker_key][p_name]["cost"] += sec_cost
        else:
            if key not in issues_secondary:
                issues_secondary[key] = {
                    "series": h.sgb_series,
                    "maturity_date": h.sgb_maturity_date,
                    "isin": h.isin,
                    "last_price": last_price,
                    "holdings": {p: {"qty": Decimal("0"), "cost": Decimal("0")} for p in PERSONS}
                }
            issues_secondary[key]["holdings"][p_name]["qty"] += qty
            issues_secondary[key]["holdings"][p_name]["cost"] += total_cost
            broker_secondary[broker_key][p_name]["qty"] += qty
            broker_secondary[broker_key][p_name]["cost"] += total_cost
            
    # Convert dicts to sorted lists
    def to_sorted_list(issues_dict):
        result_list = []
        for key, data in sorted(issues_dict.items(), key=lambda x: x[0][2] if x[0][2] else datetime.date.max):
            row = {
                "series": data["series"],
                "maturity_date": data["maturity_date"].strftime("%d-%b-%Y") if data["maturity_date"] else "—",
                "isin": data["isin"],
                "last_price": float(data["last_price"]),
                "holdings": {}
            }
            total_qty = Decimal("0")
            total_cost = Decimal("0")
            
            for p in PERSONS:
                h_p = data["holdings"][p]
                qty_p = h_p["qty"]
                cost_p = h_p["cost"]
                avg_p = cost_p / qty_p if qty_p > 0 else Decimal("0")
                
                total_qty += qty_p
                total_cost += cost_p
                
                row["holdings"][p] = {
                    "qty": float(qty_p),
                    "avg": float(avg_p),
                    "cost": float(cost_p)
                }
                
            curr_val = total_qty * data["last_price"]
            pnl = curr_val - total_cost
            return_pct = pnl / total_cost if total_cost > 0 else Decimal("0")
            
            row["total_qty"] = float(total_qty)
            row["total_cost"] = float(total_cost)
            row["current_value"] = float(curr_val)
            row["pnl"] = float(pnl)
            row["return_pct"] = float(return_pct)
            result_list.append(row)
        return result_list
        
    consolidated_rows = to_sorted_list(issues_consolidated)
    primary_rows = to_sorted_list(issues_primary)
    secondary_rows = to_sorted_list(issues_secondary)
    
    # Process broker maps
    def process_broker_map(b_map):
        processed = {}
        for broker, p_data in b_map.items():
            processed[broker] = {}
            total_b_qty = Decimal("0")
            total_b_cost = Decimal("0")
            for p in PERSONS:
                qty = p_data[p]["qty"]
                cost = p_data[p]["cost"]
                avg = cost / qty if qty > 0 else Decimal("0")
                processed[broker][p] = {
                    "qty": float(qty),
                    "avg": float(avg),
                    "cost": float(cost)
                }
                total_b_qty += qty
                total_b_cost += cost
            b_avg = total_b_cost / total_b_qty if total_b_qty > 0 else Decimal("0")
            processed[broker]["TOTAL"] = {
                "qty": float(total_b_qty),
                "cost": float(total_b_cost),
                "avg": float(b_avg)
            }
        return processed
        
    broker_primary_proc = process_broker_map(broker_primary)
    broker_secondary_proc = process_broker_map(broker_secondary)
    
    # Calculate grand totals
    grand_totals = {
        "consolidated": {p: {"qty": 0.0, "cost": 0.0, "avg": 0.0} for p in PERSONS + ["TOTAL"]},
        "primary": {p: {"qty": 0.0, "cost": 0.0, "avg": 0.0} for p in PERSONS + ["TOTAL"]},
        "secondary": {p: {"qty": 0.0, "cost": 0.0, "avg": 0.0} for p in PERSONS + ["TOTAL"]}
    }
    
    def calc_table_totals(rows, table_key):
        total_qty = {p: Decimal("0") for p in PERSONS}
        total_cost = {p: Decimal("0") for p in PERSONS}
        
        for r in rows:
            for p in PERSONS:
                qty_val = Decimal(str(r["holdings"][p]["qty"]))
                cost_val = Decimal(str(r["holdings"][p]["cost"]))
                total_qty[p] += qty_val
                total_cost[p] += cost_val
                
        grand_q = Decimal("0")
        grand_c = Decimal("0")
        for p in PERSONS:
            q = total_qty[p]
            c = total_cost[p]
            grand_q += q
            grand_c += c
            avg = c / q if q > 0 else Decimal("0")
            grand_totals[table_key][p] = {
                "qty": float(q),
                "cost": float(c),
                "avg": float(avg)
            }
        grand_avg = grand_c / grand_q if grand_q > 0 else Decimal("0")
        grand_totals[table_key]["TOTAL"] = {
            "qty": float(grand_q),
            "cost": float(grand_c),
            "avg": float(grand_avg)
        }
        
    calc_table_totals(consolidated_rows, "consolidated")
    calc_table_totals(primary_rows, "primary")
    calc_table_totals(secondary_rows, "secondary")
    
    return {
        "as_of_date": latest_date,
        "consolidated": consolidated_rows,
        "primary": primary_rows,
        "secondary": secondary_rows,
        "broker_primary": broker_primary_proc,
        "broker_secondary": broker_secondary_proc,
        "grand_totals": grand_totals
    }

@router.get("/", response_class=HTMLResponse)
def gold_dashboard(request: Request, db: Session = Depends(get_db), success_msg: str = None, error_msg: str = None):
    # Run SGB table compiler
    tables_data = compile_sgb_tables(db)
    
    # Run flat list retrieval for the 4th tab
    holdings_list = get_flat_holdings_list(db)
        
    total_cost = Decimal("0")
    total_value = Decimal("0")
    for r in tables_data["consolidated"]:
        total_cost += Decimal(str(r["total_cost"]))
        total_value += Decimal(str(r["current_value"]))
        
    total_pnl = total_value - total_cost
    total_returns_pct = total_pnl / total_cost if total_cost > 0 else Decimal("0")
    
    summary = {
        "total_cost": float(total_cost),
        "total_value": float(total_value),
        "total_pnl": float(total_pnl),
        "total_returns_pct": float(total_returns_pct),
        "count": len(tables_data["consolidated"])
    }
    
    upload_messages = []
    if success_msg:
        upload_messages.append({"type": "success", "text": success_msg})
    if error_msg:
        upload_messages.append({"type": "error", "text": error_msg})
        
    return templates.TemplateResponse(
        request=request,
        name="gold/gold_dashboard.html",
        context={
            "as_of_date": tables_data["as_of_date"].strftime("%d-%b-%Y") if tables_data["as_of_date"] else "—",
            "consolidated": tables_data["consolidated"],
            "primary": tables_data["primary"],
            "secondary": tables_data["secondary"],
            "broker_primary": tables_data["broker_primary"],
            "broker_secondary": tables_data["broker_secondary"],
            "grand_totals": tables_data["grand_totals"],
            "holdings": holdings_list,
            "summary": summary,
            "upload_messages": upload_messages
        }
    )

# get_flat_holdings_list helper defined above

def add_pivoted_sheet(wb, sheet_name, title_text, rows, grand_totals, cmp_val):
    ws = wb.create_sheet(title=sheet_name)
    ws.views.sheetView[0].showGridLines = True
    
    # Styles
    font_title = Font(name="Calibri", size=18, bold=True, color="000000")
    font_subtitle = Font(name="Calibri", size=14, bold=True, color="555555")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Arial", size=10, bold=False, color="000000")
    font_total = Font(name="Arial", size=10, bold=True, color="000000")
    
    fill_header = PatternFill(fill_type='solid', start_color="1F4E78", end_color="1F4E78") # Sleek steel blue
    fill_total = PatternFill(fill_type='solid', start_color="F2F2F2", end_color="F2F2F2")
    
    thin_side = Side(style='thin', color='CCCCCC')
    double_side = Side(style='double', color='000000')
    thick_top = Side(style='medium', color='000000')
    
    # Title Block
    ws.merge_cells("A1:U1")
    t_cell = ws["A1"]
    t_cell.value = title_text
    t_cell.font = font_title
    t_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    
    ws.merge_cells("A2:U2")
    sub_cell = ws["A2"]
    sub_cell.value = "FAMILY WEALTH SGB PORTFOLIO"
    sub_cell.font = font_subtitle
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20
    
    # Headers
    h_row_1 = 4
    h_row_2 = 5
    ws.row_dimensions[h_row_1].height = 25
    ws.row_dimensions[h_row_2].height = 25
    
    # Set merge cells for headers
    ws.merge_cells(start_row=h_row_1, start_column=1, end_row=h_row_2, end_column=1) # Series
    ws.cell(row=h_row_1, column=1, value="Series").font = font_header
    ws.cell(row=h_row_1, column=1).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=h_row_1, column=1).fill = fill_header
    
    ws.merge_cells(start_row=h_row_1, start_column=2, end_row=h_row_2, end_column=2) # ISIN
    ws.cell(row=h_row_1, column=2, value="ISIN").font = font_header
    ws.cell(row=h_row_1, column=2).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=h_row_1, column=2).fill = fill_header
    
    ws.merge_cells(start_row=h_row_1, start_column=3, end_row=h_row_2, end_column=3) # Maturity
    ws.cell(row=h_row_1, column=3, value="Maturity").font = font_header
    ws.cell(row=h_row_1, column=3).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=h_row_1, column=3).fill = fill_header
    
    ws.merge_cells(start_row=h_row_1, start_column=4, end_row=h_row_2, end_column=4) # CMP
    ws.cell(row=h_row_1, column=4, value="CMP").font = font_header
    ws.cell(row=h_row_1, column=4).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=h_row_1, column=4).fill = fill_header
    
    persons = ["PATRIARCH", "MEMBER_1", "MEMBER_2", "MEMBER_3"]
    for i, p in enumerate(persons):
        start_col = 5 + i * 3
        ws.merge_cells(start_row=h_row_1, start_column=start_col, end_row=h_row_1, end_column=start_col+2)
        ws.cell(row=h_row_1, column=start_col, value=p).font = font_header
        ws.cell(row=h_row_1, column=start_col).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=h_row_1, column=start_col).fill = fill_header
        
        ws.cell(row=h_row_2, column=start_col, value="Qty").font = font_header
        ws.cell(row=h_row_2, column=start_col).fill = fill_header
        ws.cell(row=h_row_2, column=start_col+1, value="Avg").font = font_header
        ws.cell(row=h_row_2, column=start_col+1).fill = fill_header
        ws.cell(row=h_row_2, column=start_col+2, value="Value").font = font_header
        ws.cell(row=h_row_2, column=start_col+2).fill = fill_header
        
    ws.merge_cells(start_row=h_row_1, start_column=17, end_row=h_row_1, end_column=21)
    ws.cell(row=h_row_1, column=17, value="TOTAL").font = font_header
    ws.cell(row=h_row_1, column=17).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=h_row_1, column=17).fill = fill_header
    
    ws.cell(row=h_row_2, column=17, value="Qty").font = font_header
    ws.cell(row=h_row_2, column=17).fill = fill_header
    ws.cell(row=h_row_2, column=18, value="Cost").font = font_header
    ws.cell(row=h_row_2, column=18).fill = fill_header
    ws.cell(row=h_row_2, column=19, value="Value").font = font_header
    ws.cell(row=h_row_2, column=19).fill = fill_header
    ws.cell(row=h_row_2, column=20, value="P&L").font = font_header
    ws.cell(row=h_row_2, column=20).fill = fill_header
    ws.cell(row=h_row_2, column=21, value="Return").font = font_header
    ws.cell(row=h_row_2, column=21).fill = fill_header
    
    # Fill background borders for headers
    for r in [h_row_1, h_row_2]:
        for c in range(1, 22):
            cell = ws.cell(row=r, column=c)
            cell.border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
            if not cell.fill.fill_type:
                cell.fill = fill_header
                
    # Data rows
    curr_row = 6
    for row_data in rows:
        ws.row_dimensions[curr_row].height = 20
        
        ws.cell(row=curr_row, column=1, value=row_data["series"]).font = font_data
        ws.cell(row=curr_row, column=2, value=row_data["isin"]).font = font_data
        ws.cell(row=curr_row, column=3, value=row_data["maturity_date"]).font = font_data
        ws.cell(row=curr_row, column=4, value=row_data["last_price"]).font = font_data
        ws.cell(row=curr_row, column=4).number_format = '₹#,##0.00'
        
        for p_idx, p in enumerate(persons):
            c_p = row_data["holdings"][p]
            qty_val = c_p["qty"] if c_p["qty"] > 0 else None
            avg_val = c_p["avg"] if c_p["qty"] > 0 else None
            val_val = c_p["qty"] * row_data["last_price"] if c_p["qty"] > 0 else None
            
            c_col = 5 + p_idx * 3
            ws.cell(row=curr_row, column=c_col, value=qty_val).font = font_data
            ws.cell(row=curr_row, column=c_col).number_format = '#,##0'
            ws.cell(row=curr_row, column=c_col+1, value=avg_val).font = font_data
            ws.cell(row=curr_row, column=c_col+1).number_format = '₹#,##0.00'
            ws.cell(row=curr_row, column=c_col+2, value=val_val).font = font_data
            ws.cell(row=curr_row, column=c_col+2).number_format = '₹#,##0'
            
        ws.cell(row=curr_row, column=17, value=row_data["total_qty"]).font = font_data
        ws.cell(row=curr_row, column=17).number_format = '#,##0'
        ws.cell(row=curr_row, column=18, value=row_data["total_cost"]).font = font_data
        ws.cell(row=curr_row, column=18).number_format = '₹#,##0'
        ws.cell(row=curr_row, column=19, value=row_data["current_value"]).font = font_data
        ws.cell(row=curr_row, column=19).number_format = '₹#,##0'
        ws.cell(row=curr_row, column=20, value=row_data["pnl"]).font = font_data
        ws.cell(row=curr_row, column=20).number_format = '₹#,##0'
        ws.cell(row=curr_row, column=21, value=row_data["return_pct"]).font = font_data
        ws.cell(row=curr_row, column=21).number_format = '0.00%'
        
        for col in range(1, 22):
            cell = ws.cell(row=curr_row, column=col)
            cell.border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
            
        curr_row += 1
        
    # Totals Row
    ws.row_dimensions[curr_row].height = 22
    ws.cell(row=curr_row, column=1, value="TOTAL / WEIGHTED AVERAGE").font = font_total
    
    # Set borders and fills for totals row
    for col in range(1, 22):
        cell = ws.cell(row=curr_row, column=col)
        cell.fill = fill_total
        cell.border = Border(left=thin_side, right=thin_side, top=thick_top, bottom=double_side)
        
    for p_idx, p in enumerate(persons):
        col_letter_qty = get_column_letter(5 + p_idx*3)
        col_letter_avg = get_column_letter(6 + p_idx*3)
        col_letter_val = get_column_letter(7 + p_idx*3)
        
        q_cell = ws.cell(row=curr_row, column=5 + p_idx*3, value=f"=SUM({col_letter_qty}6:{col_letter_qty}{curr_row-1})")
        q_cell.font = font_total
        q_cell.number_format = '#,##0'
        
        a_cell = ws.cell(row=curr_row, column=6 + p_idx*3, value=f"=IF({col_letter_qty}{curr_row}>0, SUMPRODUCT({col_letter_qty}6:{col_letter_qty}{curr_row-1}, {col_letter_avg}6:{col_letter_avg}{curr_row-1})/{col_letter_qty}{curr_row}, 0)")
        a_cell.font = font_total
        a_cell.number_format = '₹#,##0.00'
        
        v_cell = ws.cell(row=curr_row, column=7 + p_idx*3, value=f"=SUM({col_letter_val}6:{col_letter_val}{curr_row-1})")
        v_cell.font = font_total
        v_cell.number_format = '₹#,##0'
        
    ws.cell(row=curr_row, column=17, value=f"=SUM(Q6:Q{curr_row-1})").font = font_total
    ws.cell(row=curr_row, column=17).number_format = '#,##0'
    ws.cell(row=curr_row, column=18, value=f"=SUM(R6:R{curr_row-1})").font = font_total
    ws.cell(row=curr_row, column=18).number_format = '₹#,##0'
    ws.cell(row=curr_row, column=19, value=f"=SUM(S6:S{curr_row-1})").font = font_total
    ws.cell(row=curr_row, column=19).number_format = '₹#,##0'
    ws.cell(row=curr_row, column=20, value=f"=S{curr_row}-R{curr_row}").font = font_total
    ws.cell(row=curr_row, column=20).number_format = '₹#,##0'
    ws.cell(row=curr_row, column=21, value=f"=IF(R{curr_row}>0, T{curr_row}/R{curr_row}, 0)").font = font_total
    ws.cell(row=curr_row, column=21).number_format = '0.00%'
    
    # Adjust column widths automatically
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if val.startswith('='):
                continue
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)
        
    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 16
    ws.column_dimensions['C'].width = 14

def add_flat_sheet(wb, sheet_name, title_text, headers, data_keys, rows, totals_enabled=True):
    ws = wb.create_sheet(title=sheet_name)
    ws.views.sheetView[0].showGridLines = True
    
    # Styles
    font_title = Font(name="Calibri", size=18, bold=True, color="000000")
    font_subtitle = Font(name="Calibri", size=14, bold=True, color="555555")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Arial", size=10, bold=False, color="000000")
    font_total = Font(name="Arial", size=10, bold=True, color="000000")
    
    fill_header = PatternFill(fill_type='solid', start_color="1F4E78", end_color="1F4E78")
    fill_total = PatternFill(fill_type='solid', start_color="F2F2F2", end_color="F2F2F2")
    
    thin_side = Side(style='thin', color='CCCCCC')
    double_side = Side(style='double', color='000000')
    thick_top = Side(style='medium', color='000000')
    
    num_cols = len(headers)
    end_col_letter = get_column_letter(num_cols)
    
    # Title Block
    ws.merge_cells(f"A1:{end_col_letter}1")
    t_cell = ws["A1"]
    t_cell.value = title_text
    t_cell.font = font_title
    t_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    
    ws.merge_cells(f"A2:{end_col_letter}2")
    sub_cell = ws["A2"]
    sub_cell.value = "SGB INDIVIDUAL HOLDINGS DETAIL"
    sub_cell.font = font_subtitle
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20
    
    # Headers
    h_row = 4
    ws.row_dimensions[h_row].height = 25
    for col_idx, h in enumerate(headers):
        cell = ws.cell(row=h_row, column=col_idx+1, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        
    # Data Rows
    curr_row = 5
    for row_data in rows:
        ws.row_dimensions[curr_row].height = 20
        for col_idx, key in enumerate(data_keys):
            cell = ws.cell(row=curr_row, column=col_idx+1)
            val = row_data.get(key)
            
            # Formatting and scaling for Tax & Maturity Details
            if sheet_name == "Tax & Maturity Details" and key in ["purchase_price", "current_price"]:
                try:
                    cell.value = float(val) / 1000 if val is not None else None
                except (ValueError, TypeError):
                    cell.value = val
                cell.number_format = '"₹"#,##0.0" K"'
                cell.alignment = Alignment(horizontal="right")
            elif sheet_name == "Tax & Maturity Details" and key in ["total_cost", "current_value", "pnl"]:
                try:
                    cell.value = float(val) / 100000 if val is not None else None
                except (ValueError, TypeError):
                    cell.value = val
                cell.number_format = '"₹"#,##0.00" L"'
                cell.alignment = Alignment(horizontal="right")
            else:
                if key in ["purchase_date", "maturity_date", "date_2_years_long"]:
                    cell.alignment = Alignment(horizontal="center")
                    if isinstance(val, str):
                        parsed_dt = None
                        for date_fmt in ["%d-%b-%Y", "%b-%y", "%Y-%m-%d"]:
                            try:
                                parsed_dt = datetime.datetime.strptime(val, date_fmt)
                                break
                            except ValueError:
                                continue
                        if parsed_dt:
                            cell.value = parsed_dt.date()
                            cell.number_format = 'mmm-yy'
                        else:
                            cell.value = val
                    elif isinstance(val, (datetime.date, datetime.datetime)):
                        cell.value = val
                        cell.number_format = 'mmm-yy'
                    else:
                        cell.value = val
                else:
                    cell.value = val
                    # Apply standard formats
                    if key in ["units", "qty"]:
                        cell.number_format = '#,##0.00'
                        cell.alignment = Alignment(horizontal="right")
                    elif key in ["days_to_lt"]:
                        cell.number_format = '#,##0'
                        cell.alignment = Alignment(horizontal="right")
                    elif key in ["purchase_price", "current_price"]:
                        cell.number_format = '₹#,##0.00'
                        cell.alignment = Alignment(horizontal="right")
                    elif key in ["total_cost", "current_value", "pnl"]:
                        cell.number_format = '₹#,##0'
                        cell.alignment = Alignment(horizontal="right")
                    elif key in ["return_pct"]:
                        cell.number_format = '0.00%'
                        cell.alignment = Alignment(horizontal="right")
                    elif key in ["owner", "broker", "type", "tax", "is_long_term"]:
                        cell.alignment = Alignment(horizontal="center")
                    
            cell.font = font_data
            cell.border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
            
        curr_row += 1
        
    # Totals Row
    if totals_enabled and len(rows) > 0:
        ws.row_dimensions[curr_row].height = 22
        for col_idx in range(num_cols):
            cell = ws.cell(row=curr_row, column=col_idx+1)
            cell.fill = fill_total
            cell.border = Border(left=thin_side, right=thin_side, top=thick_top, bottom=double_side)
            
        ws.cell(row=curr_row, column=1, value="TOTAL / WEIGHTED AVERAGE").font = font_total
        
        # Calculate cell formula locations based on header names
        qty_col = -1
        cost_col = -1
        val_col = -1
        pnl_col = -1
        avg_p_col = -1
        avg_c_col = -1
        ret_col = -1
        
        for col_idx, h in enumerate(headers):
            h_clean = h.lower()
            if "qty" in h_clean:
                qty_col = col_idx + 1
            elif "cost" in h_clean:
                cost_col = col_idx + 1
            elif "value" in h_clean:
                val_col = col_idx + 1
            elif "p&l" in h_clean or "profit" in h_clean:
                pnl_col = col_idx + 1
            elif "purchase price" in h_clean or "avg price" in h_clean:
                avg_p_col = col_idx + 1
            elif "current price" in h_clean:
                avg_c_col = col_idx + 1
            elif "return" in h_clean:
                ret_col = col_idx + 1
                
        if qty_col > 0:
            qty_letter = get_column_letter(qty_col)
            qty_cell = ws.cell(row=curr_row, column=qty_col, value=f"=SUM({qty_letter}5:{qty_letter}{curr_row-1})")
            qty_cell.font = font_total
            qty_cell.number_format = '#,##0.00'
            
        if cost_col > 0:
            cost_letter = get_column_letter(cost_col)
            cost_cell = ws.cell(row=curr_row, column=cost_col, value=f"=SUM({cost_letter}5:{cost_letter}{curr_row-1})")
            cost_cell.font = font_total
            if sheet_name == "Tax & Maturity Details":
                cost_cell.number_format = '"₹"#,##0.00" L"'
            else:
                cost_cell.number_format = '₹#,##0'
            
        if val_col > 0:
            val_letter = get_column_letter(val_col)
            val_cell = ws.cell(row=curr_row, column=val_col, value=f"=SUM({val_letter}5:{val_letter}{curr_row-1})")
            val_cell.font = font_total
            if sheet_name == "Tax & Maturity Details":
                val_cell.number_format = '"₹"#,##0.00" L"'
            else:
                val_cell.number_format = '₹#,##0'
            
        if pnl_col > 0:
            pnl_letter = get_column_letter(pnl_col)
            pnl_cell = ws.cell(row=curr_row, column=pnl_col, value=f"=SUM({pnl_letter}5:{pnl_letter}{curr_row-1})")
            pnl_cell.font = font_total
            if sheet_name == "Tax & Maturity Details":
                pnl_cell.number_format = '"₹"#,##0.00" L"'
            else:
                pnl_cell.number_format = '₹#,##0'
            
        if avg_p_col > 0 and qty_col > 0 and cost_col > 0:
            qty_letter = get_column_letter(qty_col)
            cost_letter = get_column_letter(cost_col)
            if sheet_name == "Tax & Maturity Details":
                avg_p_cell = ws.cell(row=curr_row, column=avg_p_col, value=f"=IF({qty_letter}{curr_row}>0, ({cost_letter}{curr_row}*100)/{qty_letter}{curr_row}, 0)")
                avg_p_cell.number_format = '"₹"#,##0.0" K"'
            else:
                avg_p_cell = ws.cell(row=curr_row, column=avg_p_col, value=f"=IF({qty_letter}{curr_row}>0, {cost_letter}{curr_row}/{qty_letter}{curr_row}, 0)")
                avg_p_cell.number_format = '₹#,##0.00'
            avg_p_cell.font = font_total
            
        if avg_c_col > 0 and qty_col > 0 and val_col > 0:
            qty_letter = get_column_letter(qty_col)
            val_letter = get_column_letter(val_col)
            if sheet_name == "Tax & Maturity Details":
                avg_c_cell = ws.cell(row=curr_row, column=avg_c_col, value=f"=IF({qty_letter}{curr_row}>0, ({val_letter}{curr_row}*100)/{qty_letter}{curr_row}, 0)")
                avg_c_cell.number_format = '"₹"#,##0.0" K"'
            else:
                avg_c_cell = ws.cell(row=curr_row, column=avg_c_col, value=f"=IF({qty_letter}{curr_row}>0, {val_letter}{curr_row}/{qty_letter}{curr_row}, 0)")
                avg_c_cell.number_format = '₹#,##0.00'
            avg_c_cell.font = font_total
            
        if ret_col > 0 and cost_col > 0 and pnl_col > 0:
            cost_letter = get_column_letter(cost_col)
            pnl_letter = get_column_letter(pnl_col)
            ret_cell = ws.cell(row=curr_row, column=ret_col, value=f"=IF({cost_letter}{curr_row}>0, {pnl_letter}{curr_row}/{cost_letter}{curr_row}, 0)")
            ret_cell.font = font_total
            ret_cell.number_format = '0.00%'
            
    # Adjust column widths automatically
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if val.startswith('='):
                continue
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

@router.get("/export-sgb")
def export_sgb_schedule(db: Session = Depends(get_db)):
    # Run migrations and compile SGB tables
    tables_data = compile_sgb_tables(db)
    primary_rows = tables_data["primary"]
    secondary_rows = tables_data["secondary"]
    consolidated_rows = tables_data["consolidated"]
    as_of_date = tables_data["as_of_date"]
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Standard Schedule"
    ws.views.sheetView[0].showGridLines = True
    
    # Styles
    font_title = Font(name="Calibri", size=20, bold=True, color="000000", underline="single")
    font_subtitle = Font(name="Calibri", size=16, bold=True, color="000000", underline="single")
    font_header_grp = Font(name="Calibri", size=12, bold=True, color="000000")
    font_header_name = Font(name="Calibri", size=12, bold=True, color="000000")
    font_data = Font(name="Arial", size=11, bold=False, color="000000")
    font_total_label = Font(name="Arial", size=11, bold=True, color="000000")
    font_total_val = Font(name="Arial", size=11, bold=True, color="000000")
    font_invest_label = Font(name="Calibri", size=14, bold=True, color="000000", underline="single")
    font_invest_val = Font(name="Calibri", size=14, bold=False, color="000000")
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    
    # Border sides
    thin_side = Side(style='thin', color='A0A0A0')
    double_side = Side(style='double', color='000000')
    medium_side = Side(style='medium', color='000000')
    
    # Fills
    gray_fill = PatternFill(fill_type='solid', start_color="D9D9D9", end_color="D9D9D9")
    white_fill = PatternFill(fill_type=None)
    
    # Set explicit column widths to match sizes exactly
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 12
    for col_idx in [3, 5, 7, 9]:  # C, E, G, I (QTY)
        ws.column_dimensions[get_column_letter(col_idx)].width = 8.5
    for col_idx in [4, 6, 8, 10]:  # D, F, H, J (AVG)
        ws.column_dimensions[get_column_letter(col_idx)].width = 12
        
    def write_table_block(start_row, title_text, rows, cmp_val):
        # Explicit row heights
        ws.row_dimensions[start_row].height = 25.8       # Title
        ws.row_dimensions[start_row+1].height = 21.0     # Subtitle
        ws.row_dimensions[start_row+2].height = 15.0     # Spacer
        ws.row_dimensions[start_row+3].height = 23.4     # Header Grp
        ws.row_dimensions[start_row+4].height = 31.8     # Header Names
        
        # 1. Title
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=10)
        t_cell = ws.cell(row=start_row, column=1, value=title_text)
        t_cell.font = font_title
        t_cell.alignment = align_center
        
        # 2. Subtitle
        ws.merge_cells(start_row=start_row+1, start_column=1, end_row=start_row+1, end_column=10)
        sub_cell = ws.cell(row=start_row+1, column=1, value="FAMILY TRUST")
        sub_cell.font = font_subtitle
        sub_cell.alignment = align_center
        
        # 3. Headers Group
        h_grp_row = start_row + 3
        ws.merge_cells(start_row=h_grp_row, start_column=1, end_row=h_grp_row, end_column=2)
        cell = ws.cell(row=h_grp_row, column=1, value="SERIES")
        cell.font = font_header_grp
        cell.alignment = align_center
        
        cols_mapping = {3: "PATRIARCH", 5: "MEMBER_1", 7: "MEMBER_2", 9: "MEMBER_3"}
        for col_idx, name in cols_mapping.items():
            ws.merge_cells(start_row=h_grp_row, start_column=col_idx, end_row=h_grp_row, end_column=col_idx+1)
            cell = ws.cell(row=h_grp_row, column=col_idx, value=name)
            cell.font = font_header_grp
            cell.alignment = align_center
            
        for col in range(1, 11):
            cell = ws.cell(row=h_grp_row, column=col)
            l_side = double_side if col in [1, 3, 5, 7, 9] else thin_side
            r_side = double_side if col in [2, 4, 6, 8, 10] else thin_side
            if col == 1:
                l_side = medium_side
            if col == 10:
                r_side = medium_side
            cell.border = Border(left=l_side, right=r_side, top=double_side, bottom=thin_side)
            
        # 4. Header Names
        h_names_row = start_row + 4
        headers = ["ISSUE YEAR", "REDM. YEAR", "QTY", "AVG", "QTY", "AVG", "QTY", "AVG", "QTY", "AVG"]
        for col_idx, h in enumerate(headers):
            cell = ws.cell(row=h_names_row, column=col_idx+1, value=h)
            cell.font = font_header_name
            cell.alignment = align_center
            
            l_side = double_side if (col_idx+1) in [1, 3, 5, 7, 9] else thin_side
            r_side = double_side if (col_idx+1) in [2, 4, 6, 8, 10] else thin_side
            if (col_idx+1) == 1:
                l_side = medium_side
            if (col_idx+1) == 10:
                r_side = medium_side
            cell.border = Border(left=l_side, right=r_side, top=thin_side, bottom=medium_side)
            
        # 5. Data Rows
        data_start_row = start_row + 5
        current_row = data_start_row
        
        for r_idx, row_data in enumerate(rows):
            ws.row_dimensions[current_row].height = 23.4  # Explicit height for data rows
            series_name = row_data["series"] or ""
            m_year = ""
            m = re.search(r"(\d{4}-\d{2,4})", series_name)
            if m:
                m_year = m.group(1)
                
            m_date = datetime.datetime.strptime(row_data["maturity_date"], "%d-%b-%Y").date() if row_data["maturity_date"] != "—" else None
            
            # Col A (ISSUE YEAR)
            cell_a = ws.cell(row=current_row, column=1, value=m_year)
            cell_a.alignment = align_center
            cell_a.font = font_data
            
            # Col B (REDM. YEAR)
            cell_b = ws.cell(row=current_row, column=2, value=m_date)
            cell_b.number_format = 'mmm-yy'
            cell_b.alignment = align_center
            cell_b.font = font_data
            
            for p_idx, p in enumerate(PERSONS):
                c_p = row_data["holdings"][p]
                q_val = c_p["qty"] if c_p["qty"] > 0 else None
                a_val = c_p["avg"] if c_p["qty"] > 0 else None
                
                cell_q = ws.cell(row=current_row, column=3 + p_idx*2, value=q_val)
                cell_q.number_format = '#,##0'
                cell_q.alignment = align_center
                cell_q.font = font_data
                
                cell_a = ws.cell(row=current_row, column=4 + p_idx*2, value=a_val)
                cell_a.number_format = '#,##0'
                cell_a.alignment = align_center
                cell_a.font = font_data
                
            # Borders and fills
            for col in range(1, 11):
                cell = ws.cell(row=current_row, column=col)
                
                # Fills: White-gray pattern (white for A, C, D, I, J; gray for B, E, F, G, H)
                if col in [2, 5, 6, 7, 8]:
                    cell.fill = gray_fill
                else:
                    cell.fill = white_fill
                    
                l_side = double_side if col in [1, 3, 5, 7, 9] else thin_side
                r_side = double_side if col in [2, 4, 6, 8, 10] else thin_side
                if col == 1:
                    l_side = medium_side
                if col == 10:
                    r_side = medium_side
                t_b = medium_side if current_row == data_start_row else thin_side
                cell.border = Border(left=l_side, right=r_side, top=t_b, bottom=thin_side)
                
            current_row += 1
            
        data_end_row = current_row - 1
        
        # Merge issue years in Col A
        s_row = data_start_row
        curr_year = None
        for r in range(data_start_row, data_end_row + 1):
            val = ws.cell(row=r, column=1).value
            if val != curr_year:
                if curr_year is not None and r - 1 > s_row:
                    ws.merge_cells(start_row=s_row, start_column=1, end_row=r-1, end_column=1)
                    ws.cell(row=s_row, column=1).alignment = align_center
                s_row = r
                curr_year = val
        if curr_year is not None and data_end_row > s_row:
            ws.merge_cells(start_row=s_row, start_column=1, end_row=data_end_row, end_column=1)
            ws.cell(row=s_row, column=1).alignment = align_center
            
        # Last data row's bottom border is medium
        for col in range(1, 11):
            cell = ws.cell(row=data_end_row, column=col)
            cell.border = Border(left=cell.border.left, right=cell.border.right, top=cell.border.top, bottom=medium_side)
            
        # Sums Row
        tot_row = current_row
        ws.row_dimensions[tot_row].height = 23.4
        ws.cell(row=tot_row, column=1, value=None)
        ws.cell(row=tot_row, column=2, value=None)
        
        for p_idx, p in enumerate(PERSONS):
            col_letter_q = get_column_letter(3 + p_idx*2)
            col_letter_a = get_column_letter(4 + p_idx*2)
            
            q_cell = ws.cell(row=tot_row, column=3 + p_idx*2)
            q_cell.value = f"=SUM({col_letter_q}{data_start_row}:{col_letter_q}{data_end_row})"
            q_cell.number_format = '#,##0'
            q_cell.font = font_total_val
            q_cell.alignment = align_center
            
            a_cell = ws.cell(row=tot_row, column=4 + p_idx*2)
            a_cell.value = f"=IF({col_letter_q}{tot_row}>0, SUMPRODUCT({col_letter_q}{data_start_row}:{col_letter_q}{data_end_row}, {col_letter_a}{data_start_row}:{col_letter_a}{data_end_row})/{col_letter_q}{tot_row}, 0)"
            a_cell.number_format = '#,##0'
            a_cell.font = font_total_val
            a_cell.alignment = align_center
            
        for col in range(1, 11):
            cell = ws.cell(row=tot_row, column=col)
            l_side = double_side if col in [1, 3, 5, 7, 9] else thin_side
            r_side = double_side if col in [2, 4, 6, 8, 10] else thin_side
            if col == 1:
                l_side = medium_side
            if col == 10:
                r_side = medium_side
            cell.border = Border(left=l_side, right=r_side, top=medium_side, bottom=medium_side)
            
        current_row += 1
        
        # TOTAL (Cr) Row
        cr_row = current_row
        ws.row_dimensions[cr_row].height = 23.4
        ws.merge_cells(start_row=cr_row, start_column=1, end_row=cr_row, end_column=2)
        total_lbl_cell = ws.cell(row=cr_row, column=1, value="TOTAL")
        total_lbl_cell.font = font_total_label
        total_lbl_cell.alignment = align_left
        
        for p_idx in range(len(PERSONS)):
            col_idx_q = 3 + p_idx*2
            col_letter_q = get_column_letter(col_idx_q)
            col_letter_a = get_column_letter(col_idx_q + 1)
            
            ws.merge_cells(start_row=cr_row, start_column=col_idx_q, end_row=cr_row, end_column=col_idx_q+1)
            val_cell = ws.cell(row=cr_row, column=col_idx_q)
            val_cell.value = f'=ROUND(({col_letter_q}{tot_row}*{col_letter_a}{tot_row})/10000000,2) & " Cr"'
            val_cell.font = font_total_val
            val_cell.alignment = align_center
            
        for col in range(1, 11):
            cell = ws.cell(row=cr_row, column=col)
            l_side = double_side if col in [1, 3, 5, 7, 9] else None
            r_side = double_side if col in [2, 4, 6, 8, 10] else None
            if col in [4, 6, 8, 10]:
                l_side = None
            if col in [3, 5, 7, 9]:
                r_side = None
            if col == 1:
                l_side = double_side
            if col == 2:
                r_side = double_side
            if col == 10:
                r_side = double_side
            cell.border = Border(left=l_side, right=r_side, top=medium_side, bottom=double_side)
            
        current_row += 1
        
        # Invested Value & Current Value Row
        inv_row = current_row
        ws.row_dimensions[inv_row].height = 23.4
        
        ws.merge_cells(start_row=inv_row, start_column=2, end_row=inv_row, end_column=3)
        inv_lbl_cell = ws.cell(row=inv_row, column=2, value="Invested Value")
        inv_lbl_cell.font = font_invest_label
        inv_lbl_cell.alignment = align_center
        
        ws.merge_cells(start_row=inv_row, start_column=4, end_row=inv_row, end_column=5)
        inv_val_cell = ws.cell(row=inv_row, column=4)
        inv_val_cell.value = f'=ROUND(VALUE(SUBSTITUTE(C{cr_row}," Cr","")) + VALUE(SUBSTITUTE(E{cr_row}," Cr","")), 2) + ROUND(VALUE(SUBSTITUTE(G{cr_row}," Cr","")) + VALUE(SUBSTITUTE(I{cr_row}," Cr","")), 2) & " Cr"'
        inv_val_cell.font = font_invest_val
        inv_val_cell.alignment = align_center
        
        ws.merge_cells(start_row=inv_row, start_column=7, end_row=inv_row, end_column=8)
        cur_lbl_cell = ws.cell(row=inv_row, column=7, value="Current Value")
        cur_lbl_cell.font = font_invest_label
        cur_lbl_cell.alignment = align_center
        
        ws.merge_cells(start_row=inv_row, start_column=9, end_row=inv_row, end_column=10)
        cur_val_cell = ws.cell(row=inv_row, column=9)
        cur_val_cell.value = f'=ROUND((C{tot_row}+E{tot_row}+G{tot_row}+I{tot_row})*{cmp_val}/10000000,2) & " Cr"'
        cur_val_cell.font = font_invest_val
        cur_val_cell.alignment = align_center
        
        ws.cell(row=inv_row, column=2).border = Border(top=double_side)
        ws.cell(row=inv_row, column=3).border = Border(top=double_side)
        ws.cell(row=inv_row, column=7).border = Border(top=double_side)
        ws.cell(row=inv_row, column=8).border = Border(top=double_side)
        
        ws.cell(row=inv_row, column=4).border = Border(left=double_side, right=double_side, top=medium_side, bottom=double_side)
        ws.cell(row=inv_row, column=5).border = Border(right=double_side, top=medium_side, bottom=double_side)
        
        ws.cell(row=inv_row, column=9).border = Border(left=double_side, right=double_side, top=medium_side, bottom=double_side)
        ws.cell(row=inv_row, column=10).border = Border(right=double_side, top=medium_side, bottom=double_side)
        
        return inv_row
        
    sec_cmp_val = float(secondary_rows[0]["last_price"]) if secondary_rows else (float(primary_rows[0]["last_price"]) if primary_rows else 7200.0)
    sec_end_row = write_table_block(1, "SG BONDs - Secondary", secondary_rows, sec_cmp_val)
    
    # 3 blank rows spacer, then primary table block
    prim_start_row = sec_end_row + 4
    prim_cmp_val = float(primary_rows[0]["last_price"]) if primary_rows else sec_cmp_val
    write_table_block(prim_start_row, "SG BONDs - Primary", primary_rows, prim_cmp_val)
    
    # 1. Consolidated Portfolio
    consolidated_grand_totals = tables_data["grand_totals"]["consolidated"]
    consolidated_cmp = float(consolidated_rows[0]["last_price"]) if consolidated_rows else 7200.0
    add_pivoted_sheet(
        wb,
        "Consolidated Portfolio",
        "Sovereign Gold Bonds - Consolidated SGB Portfolio",
        consolidated_rows,
        consolidated_grand_totals,
        consolidated_cmp
    )
    
    # 2. Primary Market
    primary_grand_totals = tables_data["grand_totals"]["primary"]
    primary_cmp = float(primary_rows[0]["last_price"]) if primary_rows else 7200.0
    add_pivoted_sheet(
        wb,
        "Primary Market",
        "Sovereign Gold Bonds - Primary Market SGB Portfolio",
        primary_rows,
        primary_grand_totals,
        primary_cmp
    )
    
    # 3. Secondary Market
    secondary_grand_totals = tables_data["grand_totals"]["secondary"]
    secondary_cmp = float(secondary_rows[0]["last_price"]) if secondary_rows else 7200.0
    add_pivoted_sheet(
        wb,
        "Secondary Market",
        "Sovereign Gold Bonds - Secondary Market SGB Portfolio",
        secondary_rows,
        secondary_grand_totals,
        secondary_cmp
    )
    
    # Get Flat holdings data
    holdings_list = get_flat_holdings_list(db)
    
    # 4. All SGB Holdings (Flat List)
    flat_headers = [
        "Owner", "Broker", "SGB Series Name", "ISIN", "Maturity Date",
        "Qty (Grams)", "Avg Price / Gram", "Total Cost", "Current Price",
        "Current Value", "P&L", "Return %"
    ]
    flat_keys = [
        "owner", "broker", "series", "isin", "maturity_date",
        "units", "purchase_price", "total_cost", "current_price",
        "current_value", "pnl", "return_pct"
    ]
    add_flat_sheet(
        wb,
        "All SGB Holdings",
        "Sovereign Gold Bonds - All SGB Holdings (Flat List)",
        flat_headers,
        flat_keys,
        holdings_list,
        totals_enabled=True
    )
    
    # 5. Tax & Maturity Details
    detailed_headers = [
        "SGB Series", "A/C", "Broker", "Tax", "Qty (Grams)",
        "Purchase Date", "Maturity Date", "Long Term", "Days to LT",
        "Purchase Price", "Current Price", "Cost", "Total Value", "Profit"
    ]
    detailed_keys = [
        "series", "owner", "broker", "tax", "units",
        "purchase_date", "maturity_date", "is_long_term", "days_to_lt",
        "purchase_price", "current_price", "total_cost", "current_value", "pnl"
    ]
    add_flat_sheet(
        wb,
        "Tax & Maturity Details",
        "Sovereign Gold Bonds - Purchase, Maturity & Tax Details",
        detailed_headers,
        detailed_keys,
        holdings_list,
        totals_enabled=True
    )
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"SGB_Portfolio_As_On_{as_of_date.strftime('%Y%m%d') if as_of_date else 'Latest'}.xlsx"
    headers = {
        'Content-Disposition': f'attachment; filename="{filename}"'
    }
    return StreamingResponse(
        output,
        headers=headers,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

@router.post("/upload")
def upload_positions_report(
    request: Request,
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    """Processes uploaded Global Positions Report excel file."""
    if not file.filename:
        return RedirectResponse("/gold/?error_msg=No file uploaded", status_code=303)
        
    try:
        # Save uploaded file temporarily
        temp_dir = tempfile.mkdtemp()
        temp_file_path = os.path.join(temp_dir, file.filename)
        
        with open(temp_file_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
            
        # Run import
        result = import_sgb_from_excel(db, temp_file_path)
        
        # Clean up temp file
        try:
            shutil.rmtree(temp_dir)
        except Exception as rmtree_err:
            print(f"Warning: Failed to delete temp directory {temp_dir}: {rmtree_err}")
            
        if result["success"]:
            msg = f"Successfully parsed {file.filename} (As of {result['as_of_date']})! Imported {result['imported_count']} SGB holdings."
            if result.get("errors_occurred"):
                msg += f" Note: {len(result['errors_occurred'])} warnings/errors occurred during import."
            return RedirectResponse(f"/gold/?success_msg={msg}", status_code=303)
        else:
            return RedirectResponse(f"/gold/?error_msg=Failed to parse file: {result['error']}", status_code=303)
            
    except Exception as e:
        import traceback
        traceback.print_exc()
        return RedirectResponse(f"/gold/?error_msg=Error processing file: {str(e)}", status_code=303)
