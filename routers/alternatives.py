# routers/alternatives.py
from fastapi import APIRouter, Depends, Request, Form
from fastapi.responses import HTMLResponse, RedirectResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy.orm import Session
from database import get_db
from models.alternatives import Alternative, UnlistedShare, AlternativeTransaction, AlternativeValuation, UnlistedTransaction
from models.persons import Person
from models.accounts import Account
from decimal import Decimal
from datetime import date, datetime
from sqlalchemy import func
import io

router = APIRouter(prefix="/alternatives", tags=["alternatives"])
templates = Jinja2Templates(directory="templates")

# Register the format_inr filter (reuse from mutual_funds router)
from routers.mutual_funds import format_inr
templates.env.filters["format_inr"] = format_inr

from services.aif_sync_engine import sync_aif_from_excel, get_sync_status, auto_sync_if_needed


# ──────────────────────────────────────────────
# UNLISTED SHARES DASHBOARD
# ──────────────────────────────────────────────

@router.get("/unlisted/", response_class=HTMLResponse)
def unlisted_dashboard(request: Request, db: Session = Depends(get_db)):
    shares = db.query(UnlistedShare).order_by(UnlistedShare.company_name).all()

    # Build person and account caches
    person_ids = list({s.person_id for s in shares})
    persons = {p.id: p for p in db.query(Person).filter(Person.id.in_(person_ids)).all()} if person_ids else {}

    account_ids = list({s.account_id for s in shares if s.account_id})
    accounts = {a.id: a for a in db.query(Account).filter(Account.id.in_(account_ids)).all()} if account_ids else {}

    # Calculate total invested across ALL holdings to compute exposure
    total_invested = Decimal("0")
    total_current_value = Decimal("0")
    for s in shares:
        total_invested += s.total_cost or Decimal("0")
        total_current_value += s.current_value or Decimal("0")

    total_pnl = total_current_value - total_invested
    total_return_pct = total_pnl / total_invested if total_invested != 0 else 0

    # Group by company name for Summary and bifurcated Trades
    company_groups = {}
    for s in shares:
        comp = s.company_name
        if comp not in company_groups:
            company_groups[comp] = {
                "company_name": comp,
                "quantity": Decimal("0"),
                "total_cost": Decimal("0"),
                "current_value": Decimal("0"),
                "current_price": s.current_price or Decimal("0"),
                "lots": []
            }
        
        person = persons.get(s.person_id)
        account = accounts.get(s.account_id) if s.account_id else None
        owner = person.short_name if person else "—"
        broker = account.institution_name if account else "—"
        
        cost = s.total_cost or Decimal("0")
        curr_val = s.current_value or Decimal("0")
        pnl = curr_val - cost
        ret_pct = pnl / cost if cost != 0 else 0
        
        company_groups[comp]["quantity"] += s.quantity or Decimal("0")
        company_groups[comp]["total_cost"] += cost
        company_groups[comp]["current_value"] += curr_val
        if s.current_price:
            company_groups[comp]["current_price"] = s.current_price  # align all to latest CMP

        company_groups[comp]["lots"].append({
            "id": s.id,
            "owner": owner,
            "broker": broker,
            "quantity": float(s.quantity or 0),
            "purchase_price": float(s.purchase_price or 0),
            "total_cost": float(cost),
            "current_price": float(s.current_price or 0),
            "current_value": float(curr_val),
            "pnl": float(pnl),
            "return_pct": float(ret_pct),
        })

    # Get transaction history from database
    txns = db.query(UnlistedTransaction).order_by(UnlistedTransaction.transaction_date.desc()).all()
    company_txns = {}
    for t in txns:
        comp = t.company_name
        if comp not in company_txns:
            company_txns[comp] = []
        company_txns[comp].append({
            "id": t.id,
            "date": t.transaction_date.strftime("%d-%b-%Y"),
            "type": t.transaction_type,
            "quantity": float(t.quantity),
            "price": float(t.price or 0),
            "amount": float(t.amount),
            "broker": t.broker or "—",
            "notes": t.notes or ""
        })

    # Build summary list
    summary_list = []
    for comp, data in company_groups.items():
        qty = data["quantity"]
        cost = data["total_cost"]
        val = data["current_value"]
        cmp_val = data["current_price"]
        
        pnl = val - cost
        ret_pct = pnl / cost if cost != 0 else 0
        avg_price = cost / qty if qty != 0 else 0
        exposure_pct = cost / total_invested if total_invested != 0 else 0
        
        summary_list.append({
            "company_name": comp,
            "quantity": float(qty),
            "current_price": float(cmp_val),
            "avg_price": float(avg_price),
            "total_cost": float(cost),
            "current_value": float(val),
            "pnl": float(pnl),
            "return_pct": float(ret_pct),
            "exposure_pct": float(exposure_pct),
            "lots_count": len(data["lots"]),
            "txns_count": len(company_txns.get(comp, []))
        })

    # Overall portfolio summary
    summary = {
        "total_invested": float(total_invested),
        "total_current_value": float(total_current_value),
        "total_pnl": float(total_pnl),
        "total_return_pct": float(total_return_pct),
        "count": len(shares),
        "assets_count": len(company_groups)
    }

    # Prepare data sorted by company name
    summary_list = sorted(summary_list, key=lambda x: x["company_name"])
    
    # Trades list grouped by company
    bifurcated_trades = []
    for comp in sorted(company_groups.keys()):
        bifurcated_trades.append({
            "company_name": comp,
            "current_price": float(company_groups[comp]["current_price"]),
            "lots": company_groups[comp]["lots"],
            "transactions": company_txns.get(comp, [])
        })

    return templates.TemplateResponse(
        request=request,
        name="alternatives/unlisted_dashboard.html",
        context={
            "summary_list": summary_list,
            "bifurcated_trades": bifurcated_trades,
            "summary": summary
        },
    )


@router.post("/unlisted/update-cmp")
async def update_unlisted_cmp(request: Request, db: Session = Depends(get_db)):
    """Update CMP for all lots of a company."""
    form_data = await request.form()
    company_name = form_data.get("company_name")
    new_price = form_data.get("current_price")

    if not company_name or new_price is None:
        return {"error": "Invalid data"}

    price_val = Decimal(str(new_price).replace(",", "").strip())

    shares = db.query(UnlistedShare).filter(UnlistedShare.company_name == company_name).all()
    for s in shares:
        s.current_price = price_val
        s.current_value = s.quantity * price_val
        s.last_valuation_date = date.today()

    db.commit()

    # Check if AJAX
    accept = request.headers.get("accept", "")
    if "application/json" in accept:
        # compute new company totals
        total_qty = sum(float(s.quantity) for s in shares)
        total_cost = sum(float(s.total_cost) for s in shares)
        total_value = sum(float(s.current_value) for s in shares)
        pnl = total_value - total_cost
        ret_pct = pnl / total_cost if total_cost != 0 else 0
        return {
            "success": True,
            "current_price": float(price_val),
            "total_value": total_value,
            "pnl": pnl,
            "return_pct": ret_pct
        }

    return RedirectResponse("/alternatives/unlisted/", status_code=303)


@router.post("/unlisted/{share_id}/update")
async def update_unlisted_share(share_id: int, request: Request, db: Session = Depends(get_db)):
    """Inline update for unlisted share quantity, price, or CMP."""
    share = db.query(UnlistedShare).filter(UnlistedShare.id == share_id).first()
    if not share:
        return {"error": "Share not found"}

    form_data = await request.form()
    new_price = form_data.get("current_price")
    new_qty = form_data.get("quantity")
    purchase_price = form_data.get("purchase_price")

    if new_price is not None and str(new_price).strip():
        # Update CMP for all lots of the same company to keep it aligned
        price_val = Decimal(str(new_price).replace(",", "").strip())
        shares = db.query(UnlistedShare).filter(UnlistedShare.company_name == share.company_name).all()
        for s in shares:
            s.current_price = price_val
            s.current_value = s.quantity * price_val
            s.last_valuation_date = date.today()
    
    if new_qty is not None and str(new_qty).strip():
        share.quantity = Decimal(str(new_qty).replace(",", "").strip())
        if share.current_price:
            share.current_value = share.quantity * share.current_price
        if share.purchase_price:
            share.total_cost = share.quantity * share.purchase_price

    if purchase_price is not None and str(purchase_price).strip():
        share.purchase_price = Decimal(str(purchase_price).replace(",", "").strip())
        share.total_cost = share.quantity * share.purchase_price

    db.commit()

    # Check if it's an AJAX request
    accept = request.headers.get("accept", "")
    if "application/json" in accept:
        return {
            "success": True,
            "quantity": float(share.quantity or 0),
            "purchase_price": float(share.purchase_price or 0),
            "total_cost": float(share.total_cost or 0),
            "current_value": float(share.current_value or 0),
            "pnl": float((share.current_value or 0) - (share.total_cost or 0)),
        }

    return RedirectResponse("/alternatives/unlisted/", status_code=303)


@router.get("/unlisted/export-xlsx")
def export_unlisted_xlsx(db: Session = Depends(get_db)):
    """Export unlisted shares (Summary and Trades) to Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from fastapi.responses import StreamingResponse

    shares = db.query(UnlistedShare).order_by(UnlistedShare.company_name).all()
    person_ids = list({s.person_id for s in shares})
    persons = {p.id: p for p in db.query(Person).filter(Person.id.in_(person_ids)).all()} if person_ids else {}
    account_ids = list({s.account_id for s in shares if s.account_id})
    accounts = {a.id: a for a in db.query(Account).filter(Account.id.in_(account_ids)).all()} if account_ids else {}

    wb = Workbook()
    
    # Setup styles
    header_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="000000")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    border_thin = Side(border_style="thin", color="D9D9D9")
    grid_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    
    zebra_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    total_font = Font(name="Calibri", size=11, bold=True)
    
    # --------------------------------------------------
    # Sheet 1: Summary (Consolidated)
    # --------------------------------------------------
    ws1 = wb.active
    ws1.title = "Summary"
    ws1.views.sheetView[0].showGridLines = True
    
    ws1.row_dimensions[1].height = 28
    
    headers1 = ["S NO", "Stock", "QTY", "CMP", "AVG", "Invested Value", "CURRENT Value", "PROFIT/ LOSS", "% Return", "% Exposure"]
    ws1.append(headers1)
    for col_idx, h in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = grid_border

    # Calculate company aggregates
    total_invested = sum((s.total_cost or Decimal("0")) for s in shares)
    
    company_groups = {}
    for s in shares:
        comp = s.company_name
        if comp not in company_groups:
            company_groups[comp] = {
                "qty": Decimal("0"),
                "cost": Decimal("0"),
                "val": Decimal("0"),
                "cmp": s.current_price or Decimal("0")
            }
        company_groups[comp]["qty"] += s.quantity or Decimal("0")
        company_groups[comp]["cost"] += s.total_cost or Decimal("0")
        company_groups[comp]["val"] += s.current_value or Decimal("0")
        if s.current_price:
            company_groups[comp]["cmp"] = s.current_price

    row_idx = 2
    for s_no, (comp, data) in enumerate(sorted(company_groups.items()), 1):
        ws1.row_dimensions[row_idx].height = 20
        qty = float(data["qty"])
        cost = float(data["cost"])
        val = float(data["val"])
        cmp_val = float(data["cmp"])
        pnl = val - cost
        ret = pnl / cost if cost != 0 else 0
        avg_price = cost / qty if qty != 0 else 0
        exposure = cost / float(total_invested) if total_invested != 0 else 0
        
        row_data = [s_no, comp, qty, cmp_val, avg_price, cost, val, pnl, ret, exposure]
        ws1.append(row_data)
        
        # Apply formatting
        is_even = (row_idx % 2 == 0)
        for col_idx in range(1, 11):
            cell = ws1.cell(row=row_idx, column=col_idx)
            cell.border = grid_border
            if is_even:
                cell.fill = zebra_fill
            
            # Alignments & Formats
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center")
            elif col_idx == 2:
                cell.alignment = Alignment(horizontal="left")
            elif col_idx in [3, 4, 5]:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "#,##0.00" if col_idx in [4, 5] else "#,##0"
            elif col_idx in [6, 7, 8]:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "#,##0.00"
            elif col_idx in [9, 10]:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "0.00%"
                
        row_idx += 1

    # Total row for Summary
    ws1.row_dimensions[row_idx].height = 22
    total_val_sum = sum(float(g["val"]) for g in company_groups.values())
    total_cost_sum = float(total_invested)
    total_pnl_sum = total_val_sum - total_cost_sum
    total_ret_pct = total_pnl_sum / total_cost_sum if total_cost_sum != 0 else 0
    
    ws1.cell(row=row_idx, column=2, value="Total").font = total_font
    ws1.cell(row=row_idx, column=6, value=total_cost_sum).font = total_font
    ws1.cell(row=row_idx, column=6).number_format = "#,##0.00"
    ws1.cell(row=row_idx, column=7, value=total_val_sum).font = total_font
    ws1.cell(row=row_idx, column=7).number_format = "#,##0.00"
    ws1.cell(row=row_idx, column=8, value=total_pnl_sum).font = total_font
    ws1.cell(row=row_idx, column=8).number_format = "#,##0.00"
    ws1.cell(row=row_idx, column=9, value=total_ret_pct).font = total_font
    ws1.cell(row=row_idx, column=9).number_format = "0.00%"
    ws1.cell(row=row_idx, column=10, value=1.0).font = total_font
    ws1.cell(row=row_idx, column=10).number_format = "0.00%"

    for col_idx in range(1, 11):
        cell = ws1.cell(row=row_idx, column=col_idx)
        cell.border = Border(top=Side(style="double"), bottom=Side(style="double"), left=border_thin, right=border_thin)

    # Auto-fit columns
    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # --------------------------------------------------
    # Sheet 2: Trades (Lots)
    # --------------------------------------------------
    ws2 = wb.create_sheet(title="Trades")
    ws2.views.sheetView[0].showGridLines = True
    ws2.row_dimensions[1].height = 28
    
    headers2 = ["S NO", "Owner", "Broker", "Company Name", "Qty", "Avg Buy Price", "Total Cost", "CMP", "Current Value", "P&L", "Return %"]
    ws2.append(headers2)
    for col_idx, h in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = grid_border

    row_idx = 2
    for s_no, s in enumerate(shares, 1):
        ws2.row_dimensions[row_idx].height = 20
        person = persons.get(s.person_id)
        account = accounts.get(s.account_id)
        cost = float(s.total_cost or 0)
        curr_val = float(s.current_value or 0)
        pnl = curr_val - cost
        ret = pnl / cost if cost else 0
        
        row_data = [
            s_no,
            person.short_name if person else "",
            account.institution_name if account else "",
            s.company_name,
            float(s.quantity or 0),
            float(s.purchase_price or 0),
            cost,
            float(s.current_price or 0),
            curr_val,
            pnl,
            ret
        ]
        ws2.append(row_data)
        
        is_even = (row_idx % 2 == 0)
        for col_idx in range(1, 12):
            cell = ws2.cell(row=row_idx, column=col_idx)
            cell.border = grid_border
            if is_even:
                cell.fill = zebra_fill
            
            # Formatting
            if col_idx in [1, 2, 3]:
                cell.alignment = Alignment(horizontal="center")
            elif col_idx == 4:
                cell.alignment = Alignment(horizontal="left")
            elif col_idx in [5, 6, 8]:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "#,##0.00" if col_idx in [6, 8] else "#,##0"
            elif col_idx in [7, 9, 10]:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "#,##0.00"
            elif col_idx == 11:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "0.00%"
                
        row_idx += 1

    # Total row for Trades
    ws2.row_dimensions[row_idx].height = 22
    ws2.cell(row=row_idx, column=4, value="Total").font = total_font
    ws2.cell(row=row_idx, column=7, value=total_cost_sum).font = total_font
    ws2.cell(row=row_idx, column=7).number_format = "#,##0.00"
    ws2.cell(row=row_idx, column=9, value=total_val_sum).font = total_font
    ws2.cell(row=row_idx, column=9).number_format = "#,##0.00"
    ws2.cell(row=row_idx, column=10, value=total_pnl_sum).font = total_font
    ws2.cell(row=row_idx, column=10).number_format = "#,##0.00"
    ws2.cell(row=row_idx, column=11, value=total_ret_pct).font = total_font
    ws2.cell(row=row_idx, column=11).number_format = "0.00%"

    for col_idx in range(1, 12):
        cell = ws2.cell(row=row_idx, column=col_idx)
        cell.border = Border(top=Side(style="double"), bottom=Side(style="double"), left=border_thin, right=border_thin)

    # Auto-fit columns for sheet 2
    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=unlisted_shares.xlsx"},
    )


# ──────────────────────────────────────────────
# AIF / PMS PORTFOLIO DASHBOARD
# ──────────────────────────────────────────────

def _compute_xirr_for_fund(aif, db):
    """Compute XIRR for a single AIF/PMS fund using transaction cash flows."""
    from services.performance_engine import calculate_xirr

    txs = (
        db.query(AlternativeTransaction)
        .filter(AlternativeTransaction.alternative_id == aif.id)
        .order_by(AlternativeTransaction.transaction_date)
        .all()
    )

    dates = []
    flows = []

    for tx in txs:
        notes_upper = (tx.notes or "").upper()
        # Exclude valuation snapshots from cash flows — they are mark-to-market, not actual cash
        if "VALUATION" in notes_upper:
            continue
        dates.append(tx.transaction_date)
        if tx.transaction_type == "DRAWDOWN":
            flows.append(-float(tx.amount))
        else:
            flows.append(float(tx.amount))

    # If no transactions found, use the lump-sum invested amount on the investment date
    if not dates and aif.amount_invested and aif.investment_date:
        dates.append(aif.investment_date)
        flows.append(-float(aif.amount_invested))

    # Terminal value: current portfolio value
    if aif.current_value and float(aif.current_value) > 0:
        val_date = aif.last_valuation_date or date.today()
        dates.append(val_date)
        flows.append(float(aif.current_value))

    if not dates:
        return None

    try:
        xirr_val = calculate_xirr(dates, flows)
        return round(xirr_val * 100, 2) if xirr_val is not None else None
    except Exception:
        return None


@router.post("/aif/sync-excel")
async def trigger_aif_sync(request: Request, db: Session = Depends(get_db)):
    """Trigger manual re-synchronization of AIF & PMS data from Excel."""
    try:
        res = sync_aif_from_excel(db)
        accept = request.headers.get("accept", "")
        if "application/json" in accept:
            return {"success": True, "result": res}
        return RedirectResponse("/alternatives/aif/?synced=1", status_code=303)
    except Exception as e:
        accept = request.headers.get("accept", "")
        if "application/json" in accept:
            return {"success": False, "error": str(e)}
        return RedirectResponse(f"/alternatives/aif/?error={str(e)}", status_code=303)


@router.get("/aif/", response_class=HTMLResponse)
def aif_dashboard(request: Request, db: Session = Depends(get_db)):
    # Auto-sync if Excel file was modified
    try:
        auto_sync_if_needed(db)
    except Exception as e:
        print(f"Error auto-syncing AIF from Excel: {e}")

    sync_status = get_sync_status()

    all_aifs = db.query(Alternative).order_by(Alternative.manager_name, Alternative.scheme_name).all()

    person_ids = list({a.person_id for a in all_aifs})
    persons = {p.id: p for p in db.query(Person).filter(Person.id.in_(person_ids)).all()} if person_ids else {}

    holdings = []
    exited_holdings = []
    total_invested = Decimal("0")
    total_current_value = Decimal("0")

    for aif in all_aifs:
        person = persons.get(aif.person_id)
        owner = person.short_name if person else "—"

        invested = float(aif.amount_invested or 0)
        curr_val = float(aif.current_value or 0)
        pnl = curr_val - invested
        ret_pct = pnl / invested if invested != 0 else 0

        xirr_val = _compute_xirr_for_fund(aif, db)

        inv_date_str = aif.investment_date.strftime("%d-%b-%Y") if aif.investment_date else "—"
        as_of_str = aif.last_valuation_date.strftime("%d-%b-%Y") if aif.last_valuation_date else "—"

        is_closed = (curr_val == 0) or ("CLOSED" in (aif.notes or "").upper()) or ("EXITED" in (aif.notes or "").upper())

        item = {
            "id": aif.id,
            "owner": owner,
            "fund_name": aif.scheme_name,
            "fund_type": aif.alternative_type.value if aif.alternative_type else "AIF",
            "manager_name": aif.manager_name,
            "investment_date": inv_date_str,
            "as_of_date": as_of_str,
            "amount_invested": invested,
            "current_value": curr_val,
            "pnl": pnl,
            "return_pct": ret_pct,
            "xirr": xirr_val,
            "notes": aif.notes or "",
        }

        if is_closed:
            exited_holdings.append(item)
        else:
            total_invested += Decimal(str(invested))
            total_current_value += Decimal(str(curr_val))
            holdings.append(item)

    total_pnl = float(total_current_value - total_invested)
    total_return_pct = total_pnl / float(total_invested) if float(total_invested) != 0 else 0

    summary = {
        "total_invested": float(total_invested),
        "total_current_value": float(total_current_value),
        "total_pnl": total_pnl,
        "total_return_pct": total_return_pct,
        "count": len(holdings),
    }

    exited_summary = {
        "count": len(exited_holdings),
        "total_invested": sum(h["amount_invested"] for h in exited_holdings),
    }

    return templates.TemplateResponse(
        request=request,
        name="alternatives/aif_dashboard.html",
        context={
            "holdings": holdings,
            "exited_holdings": exited_holdings,
            "summary": summary,
            "exited_summary": exited_summary,
            "sync_status": sync_status,
        },
    )


# ──────────────────────────────────────────────
# AIF EXPORT TO EXCEL
# ──────────────────────────────────────────────

@router.get("/aif/export-xlsx")
def export_aif_xlsx(db: Session = Depends(get_db)):
    from openpyxl import Workbook
    from fastapi.responses import StreamingResponse

    aifs = db.query(Alternative).order_by(Alternative.manager_name).all()
    person_ids = list({a.person_id for a in aifs})
    persons = {p.id: p for p in db.query(Person).filter(Person.id.in_(person_ids)).all()} if person_ids else {}

    wb = Workbook()
    ws = wb.active
    ws.title = "AIF PMS Portfolio"
    headers = ["Owner", "Fund Name", "Type", "Manager", "Investment Date", "Invested", "Current Value", "P&L", "Return %", "XIRR %"]
    ws.append(headers)

    for aif in aifs:
        person = persons.get(aif.person_id)
        invested = float(aif.amount_invested or 0)
        curr_val = float(aif.current_value or 0)
        pnl = curr_val - invested
        ret = pnl / invested if invested else 0
        xirr_val = _compute_xirr_for_fund(aif, db)

        ws.append([
            person.short_name if person else "",
            aif.scheme_name,
            aif.alternative_type.value if aif.alternative_type else "",
            aif.manager_name,
            aif.investment_date.strftime("%d-%b-%Y") if aif.investment_date else "",
            invested,
            curr_val,
            pnl,
            ret,
            xirr_val / 100 if xirr_val else None,
        ])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=aif_pms_portfolio.xlsx"},
    )


# ──────────────────────────────────────────────
# AIF FUND DETAIL PAGE
# ──────────────────────────────────────────────

@router.get("/aif/{aif_id}", response_class=HTMLResponse)
def aif_detail(aif_id: int, request: Request, db: Session = Depends(get_db)):
    aif = db.query(Alternative).filter(Alternative.id == aif_id).first()
    if not aif:
        return RedirectResponse("/alternatives/aif/", status_code=303)

    person = db.query(Person).filter(Person.id == aif.person_id).first()
    owner = person.short_name if person else "—"

    invested = float(aif.amount_invested or 0)
    curr_val = float(aif.current_value or 0)
    pnl = curr_val - invested
    ret_pct = pnl / invested if invested != 0 else 0
    xirr_val = _compute_xirr_for_fund(aif, db)

    fund = {
        "id": aif.id,
        "owner": owner,
        "fund_name": aif.scheme_name,
        "fund_type": aif.alternative_type.value if aif.alternative_type else "AIF",
        "manager_name": aif.manager_name,
        "investment_date": aif.investment_date.strftime("%d-%b-%Y") if aif.investment_date else "—",
        "amount_invested": invested,
        "current_value": curr_val,
        "pnl": pnl,
        "return_pct": ret_pct,
        "xirr": xirr_val,
        "notes": aif.notes or "",
    }

    # Get transactions
    txs_raw = (
        db.query(AlternativeTransaction)
        .filter(AlternativeTransaction.alternative_id == aif_id)
        .order_by(AlternativeTransaction.transaction_date)
        .all()
    )
    transactions = [
        {
            "id": tx.id,
            "transaction_date": tx.transaction_date.strftime("%d-%b-%Y") if tx.transaction_date else "",
            "transaction_type": tx.transaction_type,
            "amount": float(tx.amount or 0),
            "nav": float(tx.nav) if tx.nav else None,
            "units": float(tx.units) if tx.units else None,
            "notes": tx.notes or "",
        }
        for tx in txs_raw
    ]

    # Get valuations
    vals_raw = (
        db.query(AlternativeValuation)
        .filter(AlternativeValuation.alternative_id == aif_id)
        .order_by(AlternativeValuation.valuation_date)
        .all()
    )

    valuations = []
    prev_val = None
    for v in vals_raw:
        amt = float(v.valuation_amount or 0)
        pct_change = None
        if prev_val is not None and prev_val != 0:
            pct_change = (amt - prev_val) / prev_val
        prev_val = amt

        valuations.append({
            "id": v.id,
            "valuation_date": v.valuation_date.strftime("%d-%b-%Y") if v.valuation_date else "",
            "valuation_amount": amt,
            "notes": v.notes or "",
            "pct_change": pct_change,
        })

    persons_list = db.query(Person).filter(Person.is_active == True).order_by(Person.short_name).all()

    return templates.TemplateResponse(
        request=request,
        name="alternatives/aif_detail.html",
        context={
            "fund": fund,
            "transactions": transactions,
            "valuations": valuations,
            "persons": persons_list,
        },
    )


# ──────────────────────────────────────────────
# AIF TRANSACTION CRUD
# ──────────────────────────────────────────────

@router.post("/aif/{aif_id}/transaction/add")
async def add_aif_transaction(
    aif_id: int,
    transaction_date: str = Form(...),
    transaction_type: str = Form(...),
    amount: str = Form(...),
    nav: str = Form(None),
    units: str = Form(None),
    notes: str = Form(None),
    db: Session = Depends(get_db),
):
    txn = AlternativeTransaction(
        alternative_id=aif_id,
        transaction_date=datetime.strptime(transaction_date, "%Y-%m-%d").date(),
        transaction_type=transaction_type.upper(),
        amount=Decimal(amount.replace(",", "").strip()),
        nav=Decimal(nav.replace(",", "").strip()) if nav and nav.strip() else None,
        units=Decimal(units.replace(",", "").strip()) if units and units.strip() else None,
        notes=notes.strip() if notes else None,
    )
    db.add(txn)
    db.commit()
    return RedirectResponse(f"/alternatives/aif/{aif_id}", status_code=303)


@router.post("/aif/transaction/{txn_id}/delete")
def delete_aif_transaction(txn_id: int, db: Session = Depends(get_db)):
    txn = db.query(AlternativeTransaction).filter(AlternativeTransaction.id == txn_id).first()
    if txn:
        aif_id = txn.alternative_id
        db.delete(txn)
        db.commit()
        return RedirectResponse(f"/alternatives/aif/{aif_id}", status_code=303)
    return RedirectResponse("/alternatives/aif/", status_code=303)


# ──────────────────────────────────────────────
# AIF VALUATION CRUD
# ──────────────────────────────────────────────

@router.post("/aif/{aif_id}/valuation/add")
async def add_aif_valuation(
    aif_id: int,
    valuation_date: str = Form(...),
    valuation_amount: str = Form(...),
    notes: str = Form(None),
    db: Session = Depends(get_db),
):
    val = AlternativeValuation(
        alternative_id=aif_id,
        valuation_date=datetime.strptime(valuation_date, "%Y-%m-%d").date(),
        valuation_amount=Decimal(valuation_amount.replace(",", "").strip()),
        notes=notes.strip() if notes else None,
    )
    db.add(val)
    db.commit()
    return RedirectResponse(f"/alternatives/aif/{aif_id}", status_code=303)


@router.post("/aif/valuation/{val_id}/delete")
def delete_aif_valuation(val_id: int, db: Session = Depends(get_db)):
    val = db.query(AlternativeValuation).filter(AlternativeValuation.id == val_id).first()
    if val:
        aif_id = val.alternative_id
        db.delete(val)
        db.commit()
        return RedirectResponse(f"/alternatives/aif/{aif_id}", status_code=303)
    return RedirectResponse("/alternatives/aif/", status_code=303)



