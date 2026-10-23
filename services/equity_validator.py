import os
import glob
import tempfile
import shutil
import openpyxl
from typing import Dict, List, Any

BROKER_FOLDER_MAP = {
    "CBL": ["CENTRUM", "CENTRUM BROKING", "CENTRUM FINVERSE"],
    "JM": ["JM FINANCIAL", "JM FINANCIAL SERVICES"],
    "NUV": ["NUVAMA", "NUVAMA WEALTH", "EDELWEISS", "EDELWEISS BROKING"]
}

OWNER_NAME_MAP = {
    "NEXUS": ["NEXUS CAPITAL", "NEXUS CAPITAL PVT", "NEXUS CAPITAL PRIVATE LIMITED", "NEXUS"],
    "PATRIARCH": ["FAMILY PATRIARCH", "FAMILY PATRIARCH", "FAMILY PATRIARCH", "PATRIARCH"],
    "MR": ["FAMILY MEMBER 1", "MEMBER_1", "MR"],
    "VG": ["FAMILY MEMBER 2", "FAMILY MEMBER 2", "FAMILY MEMBER 2", "MEMBER_2", "VG"],
    "SR": ["FAMILY MEMBER 3", "MEMBER_3", "SR"],
    "APEX": ["APEX HOLDINGS PVT", "APEX HOLDINGS", "APEX"]
}

EXPECTED_TRANSACTION_COLS = ["Date", "Trans. Type", "Asset Name", "Quantity", "Price", "Amount"]

def detect_file_metadata(file_path: str) -> Dict[str, Any]:
    """
    Analyzes an Excel file in DATA/DATA to inspect its contents, broker identity,
    owner identity, and transaction format validity using fast openpyxl streaming.
    """
    filename = os.path.basename(file_path)
    folder_name = os.path.basename(os.path.dirname(file_path)).upper()
    
    parts = filename.replace(".xlsx", "").replace(".xls", "").split("_")
    filename_broker = parts[0].upper() if len(parts) >= 1 else ""
    filename_owner = parts[1].upper() if len(parts) >= 2 else ""
    filename_fy = "_".join(parts[2:]) if len(parts) >= 3 else ""

    result = {
        "file_path": file_path,
        "filename": filename,
        "folder": folder_name,
        "filename_broker": filename_broker,
        "filename_owner": filename_owner,
        "filename_fy": filename_fy,
        "is_valid": True,
        "detected_broker": "Unknown",
        "detected_owner": "Unknown",
        "format_type": "TRANSACTION_LEDGER",
        "issues": [],
        "warnings": [],
        "row_count": 0,
        "mixed_brokers": []
    }

    tmp_path = None
    try:
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        except PermissionError:
            with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
                tmp_path = tmp.name
            shutil.copy2(file_path, tmp_path)
            wb = openpyxl.load_workbook(tmp_path, read_only=True, data_only=True)
        
        ws = wb.active
        
        top_rows = []
        header_row_idx = None
        header_cols = []
        broker_col_idx = None
        data_rows_count = 0
        unique_brokers = set()

        for idx, row in enumerate(ws.iter_rows(values_only=True)):
            row_clean = [str(c).strip() if c is not None else "" for c in row]
            
            if idx < 25:
                top_rows.append(row_clean)
                
            # Check if this row is the transaction header
            if header_row_idx is None:
                matches = sum(any(exp.lower() in c.lower() for c in row_clean) for exp in EXPECTED_TRANSACTION_COLS)
                if matches >= 3:
                    header_row_idx = idx
                    header_cols = row_clean
                    for col_i, col_name in enumerate(header_cols):
                        if "broker" in col_name.lower():
                            broker_col_idx = col_i
                            break
            else:
                # We are below header row
                first_val = row_clean[0].upper() if row_clean else ""
                if "TOTAL" in first_val or "POWERED BY" in first_val:
                    continue
                if any(c for c in row_clean if c):
                    data_rows_count += 1
                    if broker_col_idx is not None and broker_col_idx < len(row_clean):
                        b_val = row_clean[broker_col_idx].strip()
                        if b_val and b_val != "0" and b_val.lower() != "nan":
                            unique_brokers.add(b_val)
        
        wb.close()
    except Exception as e:
        result["is_valid"] = False
        result["issues"].append(f"Cannot read Excel file: {str(e)}")
        return result
    finally:
        if tmp_path and os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except Exception:
                pass

    # Inspect text dump of top rows
    text_dump = " ".join([" ".join(r) for r in top_rows]).upper()

    # 1. Detect if this is a Capital Gains Summary Report instead of Transactions
    if "CAPITAL GAIN" in text_dump and ("SHORT-TERM" in text_dump or "LONG-TERM" in text_dump or "SCHEDULE CG" in text_dump):
        if header_row_idx is None:
            result["is_valid"] = False
            result["format_type"] = "CAPITAL_GAINS_SUMMARY"
            result["issues"].append("Capital Gains Summary Sheet. Update with transaction ledger when ready.")
            return result

    # 2. Detect Broker in header/text
    detected_broker = "Unknown"
    for b_code, keywords in BROKER_FOLDER_MAP.items():
        if any(kw in text_dump for kw in keywords):
            detected_broker = b_code
            break
    result["detected_broker"] = detected_broker

    # 3. Detect Owner in header/text
    # Note: Top-right corner often contains managing company address ("APEX HOLDINGS PRIVATE LIMITED", accounts@familyoffice.internal).
    # Specific entity names (NEXUS, PATRIARCH, MR, VG, SR) are checked in priority order.
    detected_owner = "Unknown"
    for o_code in ["NEXUS", "PATRIARCH", "MR", "VG", "SR", "APEX"]:
        keywords = OWNER_NAME_MAP[o_code]
        if any(kw in text_dump for kw in keywords):
            detected_owner = o_code
            break
    result["detected_owner"] = detected_owner

    # 4. Check for Folder / Filename vs Content Broker Mismatch
    if detected_broker != "Unknown":
        if folder_name in BROKER_FOLDER_MAP and detected_broker != folder_name:
            result["warnings"].append(f"Folder is '{folder_name}' but content mentions '{detected_broker}'. (Row-level routing applies).")
        elif filename_broker in BROKER_FOLDER_MAP and detected_broker != filename_broker:
            result["warnings"].append(f"Filename is '{filename_broker}' but content mentions '{detected_broker}'. (Row-level routing applies).")

    # 5. Check for Owner Mismatch
    if detected_owner != "Unknown" and filename_owner in OWNER_NAME_MAP and detected_owner != filename_owner:
        result["warnings"].append(f"Owner Notice: Filename says '{filename_owner}' but header mentions '{detected_owner}'.")

    # 6. Check Transaction Header Row
    if header_row_idx is None:
        result["is_valid"] = False
        result["issues"].append("Transaction Header not found (missing Date, Trans. Type, Asset Name, Quantity, Price, Amount).")
        return result

    # 7. Check for Mixed Brokers in the data rows
    foreign_brokers = []
    for b_val in unique_brokers:
        b_val_upper = b_val.upper()
        for b_code, keywords in BROKER_FOLDER_MAP.items():
            if b_code != folder_name and any(kw in b_val_upper for kw in keywords):
                foreign_brokers.append(f"{b_val} (belongs to {b_code})")

    if foreign_brokers:
        result["mixed_brokers"] = foreign_brokers
        result["warnings"].append(f"Mixed Broker Rows: Contains transactions for {', '.join(set(foreign_brokers))}.")

    result["row_count"] = data_rows_count
    return result


_FILE_CACHE = {}

def get_file_cache_key(file_path: str) -> str:
    try:
        stat = os.stat(file_path)
        return f"{file_path}_{stat.st_mtime}_{stat.st_size}"
    except Exception:
        return file_path

def get_cached_file_metadata(file_path: str) -> Dict[str, Any]:
    key = get_file_cache_key(file_path)
    if key in _FILE_CACHE:
        return _FILE_CACHE[key]
    diag = detect_file_metadata(file_path)
    _FILE_CACHE[key] = diag
    return diag

def validate_all_equity_files(force_refresh: bool = False) -> Dict[str, Any]:
    """
    Scans DATA/DATA/ and validates all Excel transaction sheets with fast mtime caching.
    Returns summary statistics and file-by-file diagnostic details.
    """
    if force_refresh:
        _FILE_CACHE.clear()

    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    data_dir = os.path.join(base_dir, "DATA", "DATA")
    
    if not os.path.exists(data_dir):
        return {
            "total_files": 0,
            "valid_files": 0,
            "invalid_count": 0,
            "warning_count": 0,
            "has_errors": False,
            "files": []
        }

    xlsx_files = glob.glob(os.path.join(data_dir, "*", "*.xlsx"))
    xlsx_files = [f for f in xlsx_files if not os.path.basename(f).startswith("~$")]

    diagnostics = []
    invalid_count = 0
    warning_count = 0

    for f in sorted(xlsx_files):
        diag = get_cached_file_metadata(f)
        if not diag["is_valid"]:
            invalid_count += 1
        elif diag["warnings"]:
            warning_count += 1
        diagnostics.append(diag)

    return {
        "total_files": len(diagnostics),
        "valid_files": len(diagnostics) - invalid_count,
        "invalid_count": invalid_count,
        "warning_count": warning_count,
        "has_errors": invalid_count > 0,
        "files": diagnostics
    }
