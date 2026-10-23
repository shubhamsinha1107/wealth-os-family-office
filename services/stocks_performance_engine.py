# services/stocks_performance_engine.py
import os
import re
import shutil
import tempfile
import openpyxl
from datetime import datetime

GLOBAL_REPORT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "DATA", "GLOBAL REPORT - MONTHLY HOLDING")
FILENAME_PATTERN = re.compile(r"MProfit Global Reporting Workbook - As on (\d{2}-\d{2}-\d{4})(?:\s+.*)?\.xlsx", re.IGNORECASE)

def extract_date_from_filename(filename):
    match = FILENAME_PATTERN.match(filename)
    if match:
        date_str = match.group(1)
        try:
            return datetime.strptime(date_str, "%d-%m-%Y").date()
        except ValueError:
            return None
    return None

def scan_and_sort_files():
    if not os.path.exists(GLOBAL_REPORT_DIR):
        return []
    
    files_with_dates = []
    for f in os.listdir(GLOBAL_REPORT_DIR):
        if f.endswith(".xlsx") and not f.startswith("~$"):
            d = extract_date_from_filename(f)
            if d:
                files_with_dates.append((d, os.path.join(GLOBAL_REPORT_DIR, f)))
                
    # Sort by date ascending
    files_with_dates.sort(key=lambda x: x[0])
    return files_with_dates

def load_stocks_holdings(file_path):
    # Copy to temporary file to bypass Windows locks
    temp_path = tempfile.mktemp(suffix=".xlsx")
    shutil.copy2(file_path, temp_path)
    
    try:
        wb = openpyxl.load_workbook(temp_path, data_only=True)
        if "Raw data" not in wb.sheetnames:
            wb.close()
            return {}
        ws = wb["Raw data"]
        
        holdings = {}
        for r in range(2, ws.max_row + 1):
            ac = ws.cell(row=r, column=3).value
            name = ws.cell(row=r, column=4).value
            port = ws.cell(row=r, column=6).value
            qty = ws.cell(row=r, column=7).value or 0.0
            invested = ws.cell(row=r, column=8).value or 0.0
            current = ws.cell(row=r, column=9).value or 0.0
            price = ws.cell(row=r, column=13).value or 0.0
            isin = ws.cell(row=r, column=14).value
            
            if not ac or str(ac).strip().upper() != "STOCKS":
                continue
                
            # Classify broker
            port_upper = str(port).upper()
            if "CENTRUM" in port_upper or "CBL" in port_upper:
                broker = "CBL / Centrum"
            elif "NUVAMA" in port_upper:
                broker = "Nuvama"
            elif "JM" in port_upper:
                broker = "JM Financial"
            else:
                broker = "Unknown"
                
            key_isin = str(isin).strip() if isin else str(name).strip()
            key = (broker, key_isin)
            
            if key not in holdings:
                holdings[key] = {
                    "name": name,
                    "qty": 0.0,
                    "invested": 0.0,
                    "current": 0.0,
                    "price": 0.0
                }
                
            holdings[key]["qty"] += float(qty)
            holdings[key]["invested"] += float(invested)
            holdings[key]["current"] += float(current)
            holdings[key]["price"] = float(price)
            
        wb.close()
        return holdings
    finally:
        try:
            os.remove(temp_path)
        except:
            pass

def calculate_period_performance(start_data, end_data):
    brokers = ["CBL / Centrum", "Nuvama", "JM Financial"]
    
    report = {}
    for b in brokers:
        report[b] = {
            "opening_aum": 0.0,
            "ending_aum": 0.0,
            "net_cash_flow": 0.0,
            "organic_growth": 0.0,
            "dietz_return": 0.0
        }
        
    all_keys = set(start_data.keys()).union(set(end_data.keys()))
    
    for key in all_keys:
        broker, isin = key
        if broker not in brokers:
            continue
            
        start = start_data.get(key, {"qty": 0.0, "invested": 0.0, "current": 0.0, "price": 0.0})
        end = end_data.get(key, {"qty": 0.0, "invested": 0.0, "current": 0.0, "price": 0.0})
        
        qty_start = start["qty"]
        qty_end = end["qty"]
        inv_start = start["invested"]
        inv_end = end["invested"]
        val_start = start["current"]
        val_end = end["current"]
        price_start = start["price"]
        price_end = end["price"]
        
        # Accumulate AUMs
        report[broker]["opening_aum"] += val_start
        report[broker]["ending_aum"] += val_end
        
        # Calculate cash flow
        cf = 0.0
        if qty_start == 0.0 and qty_end > 0.0:
            cf = inv_end
        elif qty_start > 0.0 and qty_end == 0.0:
            cf = -qty_start * price_start
        elif qty_end > qty_start:
            cf = inv_end - inv_start
        elif qty_end < qty_start:
            avg_price = (price_start + price_end) / 2.0 if (price_start > 0 and price_end > 0) else (price_start or price_end)
            cf = -(qty_start - qty_end) * avg_price
            
        report[broker]["net_cash_flow"] += cf

    total_opening = 0.0
    total_ending = 0.0
    total_cf = 0.0
    
    for b in brokers:
        rep = report[b]
        rep["organic_growth"] = rep["ending_aum"] - rep["opening_aum"] - rep["net_cash_flow"]
        denom = rep["opening_aum"] + (rep["net_cash_flow"] * 0.5)
        rep["dietz_return"] = (rep["organic_growth"] / denom * 100.0) if denom > 0.001 else 0.0
        
        total_opening += rep["opening_aum"]
        total_ending += rep["ending_aum"]
        total_cf += rep["net_cash_flow"]
        
    total_organic = total_ending - total_opening - total_cf
    total_denom = total_opening + (total_cf * 0.5)
    total_dietz = (total_organic / total_denom * 100.0) if total_denom > 0.001 else 0.0
    
    report["TOTAL"] = {
        "opening_aum": total_opening,
        "ending_aum": total_ending,
        "net_cash_flow": total_cf,
        "organic_growth": total_organic,
        "dietz_return": total_dietz
    }
    
    return report

def generate_stocks_performance_history():
    files_with_dates = scan_and_sort_files()
    if len(files_with_dates) < 2:
        return []
        
    reports = []
    # Load holdings for all files
    holdings_by_date = {}
    for d, path in files_with_dates:
        holdings_by_date[d] = load_stocks_holdings(path)
        
    # Calculate performance for each consecutive pair
    for i in range(len(files_with_dates) - 1):
        d_start, path_start = files_with_dates[i]
        d_end, path_end = files_with_dates[i+1]
        
        # Period label (e.g., "July 2026" or "August 2026")
        period_label = d_end.strftime("%B %Y")
        
        perf = calculate_period_performance(holdings_by_date[d_start], holdings_by_date[d_end])
        reports.append({
            "period": period_label,
            "start_date": d_start.strftime("%Y-%m-%d"),
            "end_date": d_end.strftime("%Y-%m-%d"),
            "data": perf
        })
        
    # Return sorted descending (newest period first) for presentation
    reports.reverse()
    return reports
